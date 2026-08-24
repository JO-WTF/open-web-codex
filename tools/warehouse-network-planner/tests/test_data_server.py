from __future__ import annotations

import asyncio
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from open_web_codex_provider import ProviderContractError, ResourceStore
from openpyxl import Workbook
from pydantic import ValidationError
from supply_chain_planner.data import server as data_server
from supply_chain_planner.data.mapping import SourceRole
from supply_chain_planner.data.workspace_intake import SourceInspectionSnapshot
from supply_chain_planner.shared.models import (
    DataPreparationNeedsInput,
    DataPreparationReady,
    DataPreparationToolResult,
    DataSourceRequirement,
    PreparationRoleCounts,
    SourceInspectionIdentity,
    SourceSelection,
)
from supply_chain_planner.shared.resources import SupplyChainResources

PREPARED_OUTPUT_DIR = "outputs/warehouse-network/prepared"


def _use_store(tmp_path, monkeypatch) -> ResourceStore:
    resources = SupplyChainResources(tmp_path, tmp_path / "profile")
    monkeypatch.setattr(data_server, "_supply_chain_resources", resources)
    monkeypatch.setattr(data_server, "_workspace", lambda _ctx: tmp_path)
    return resources.store


def _context(workspace) -> SimpleNamespace:
    meta = SimpleNamespace(
        model_extra={"codex/sandbox-state-meta": {"sandboxCwd": workspace.as_uri()}}
    )
    return SimpleNamespace(request_context=SimpleNamespace(meta=meta))


def _decisions() -> list[SourceSelection]:
    return [
        SourceSelection(relative_path="demand.csv", unit_ref="table", role=SourceRole.DEMAND),
        SourceSelection(
            relative_path="warehouses.csv",
            unit_ref="table",
            role=SourceRole.EXISTING_WAREHOUSE,
        ),
    ]


def test_data_server_exposes_workspace_preparation_not_cross_agent_data_resources() -> None:
    tools = {tool.name: tool for tool in asyncio.run(data_server.mcp.list_tools())}

    assert set(tools) == {
        "discover_workspace_sources",
        "inspect_workspace_sources",
        "prepare_network_input",
        "prepare_network_geography",
    }
    assert "normalize_candidate_delta" not in tools
    assert "derive_normalized_network_input" not in tools
    assert asyncio.run(data_server.mcp.list_resources()) == []
    assert asyncio.run(data_server.mcp.list_resource_templates()) == []

    prepare = tools["prepare_network_input"]
    assert set(prepare.inputSchema["required"]) == {
        "inspection_identity",
        "inspected_relative_paths",
        "source_selections",
        "country_code",
        "output_relative_path",
    }
    assert "confirmed_sources" not in prepare.inputSchema["properties"]
    assert prepare.inputSchema["properties"]["source_selections"]["minItems"] == 1
    assert prepare.inputSchema["properties"]["source_selections"]["maxItems"] == 640
    assert prepare.outputSchema["discriminator"]["propertyName"] == "outcome"
    assert set(prepare.outputSchema["discriminator"]["mapping"]) == {
        "ready",
        "needs_input",
        "source_changed",
    }
    assert "prepared" not in prepare.outputSchema["discriminator"]["mapping"]
    needs_input_schema = prepare.outputSchema["$defs"]["DataPreparationNeedsInput"]
    assert "prepared_input_relative_path" not in needs_input_schema["properties"]
    assert "input_identity" not in needs_input_schema["properties"]
    assert {"requirement_count", "requirements_truncated"} <= set(
        needs_input_schema["properties"]
    )
    ready_schema = prepare.outputSchema["$defs"]["DataPreparationReady"]
    assert {"warning_count", "warnings_truncated"} <= set(ready_schema["properties"])
    source_changed_schema = prepare.outputSchema["$defs"]["DataPreparationSourceChanged"]
    assert set(source_changed_schema["required"]) == {
        "outcome",
        "summary",
        "next_action",
        "retryable",
    }
    assert all(
        "resource_ref" not in definition.get("properties", {})
        for definition in prepare.outputSchema["$defs"].values()
        if isinstance(definition, dict)
    )

    inspect_tool = tools["inspect_workspace_sources"]
    assert "country_code" not in inspect_tool.inputSchema["required"]
    required_role_schema = inspect_tool.inputSchema["properties"]["required_roles"]
    planning_role_definition = inspect_tool.inputSchema["$defs"][
        required_role_schema["items"]["$ref"].split("/")[-1]
    ]
    assert set(planning_role_definition["enum"]) == {
        "demand",
        "existing_warehouse",
        "candidate_warehouse",
        "current_assignment",
        "route_quote",
    }
    assert "administrative_catalog" not in planning_role_definition["enum"]
    inspection_schema = inspect_tool.outputSchema
    assert inspection_schema["discriminator"]["propertyName"] == "outcome"
    assert set(inspection_schema["discriminator"]["mapping"]) == {
        "inspected",
        "selection_required",
        "prepared_ready",
        "prepared_selection_required",
    }
    assert len(inspection_schema["oneOf"]) == 4
    branch_required = [
        set(inspection_schema["$defs"][ref["$ref"].split("/")[-1]]["required"])
        for ref in inspection_schema["oneOf"]
    ]
    assert {
        "outcome",
        "schemaVersion",
        "summary",
        "next_action",
        "retryable",
        "source_profile",
        "inspection_identity",
        "inspected_relative_paths",
        "warnings",
        "warning_count",
        "warnings_truncated",
    } in branch_required
    assert {
        "outcome",
        "schemaVersion",
        "summary",
        "code",
        "next_action",
        "retryable",
        "observed",
        "limit",
    } in branch_required

    geography = tools["prepare_network_geography"]
    assert set(geography.inputSchema["required"]) == {
        "prepared_input_relative_path",
        "administrative_catalog_relative_path",
        "output_relative_path",
    }
    assert prepare.annotations is not None
    assert prepare.annotations.model_dump(by_alias=True, exclude_none=True) == {
        "readOnlyHint": False,
        "destructiveHint": False,
        "idempotentHint": False,
        "openWorldHint": False,
    }


def test_missing_warehouse_type_is_typed_non_retryable_user_input(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity\nCITY-1,Jakarta,50\n",
        encoding="utf-8",
    )
    (tmp_path / "warehouses.csv").write_text(
        "warehouse_id,warehouse_name,city_id,city_name,is_existing\n"
        "WH-1,Jakarta Center,CITY-1,Jakarta,true\n",
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    inspected = data_server.inspect_workspace_sources(
        ["demand.csv", "warehouses.csv"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    )
    assert inspected.structuredContent is not None
    inspection = inspected.structuredContent

    assert inspection["outcome"] == "inspected"
    assert inspection["schemaVersion"] == "workspace_source_profile.v3"
    assert inspection["next_action"] == "confirm_sources"
    assert inspection["retryable"] is False
    warehouse_profile = next(
        source
        for source in inspection["source_profile"]["sources"]
        if source["relative_path"] == "warehouses.csv"
    )
    assert "role_assessments" not in warehouse_profile
    assert len(warehouse_profile["units"]) == 1
    warehouse_unit = warehouse_profile["units"][0]
    assessment = warehouse_unit["role_assessments"][0]
    assert assessment["role"] == "existing_warehouse"
    assert assessment["state"] == "partial"
    assert assessment["missing_required_fields"] == ["warehouse_type"]
    assert {item["target_field"] for item in assessment["candidate_mappings"]} == {
        "city_id",
        "city_name",
        "is_existing",
        "warehouse_id",
        "warehouse_name",
    }
    assert [item["name"] for item in warehouse_unit["fields"]] == [
        "warehouse_id",
        "warehouse_name",
        "city_id",
        "city_name",
        "is_existing",
    ]
    assert "preview" not in warehouse_unit

    output_path = f"{PREPARED_OUTPUT_DIR}/must-not-write.json"
    blocked = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(inspection["inspection_identity"]),
        inspection["inspected_relative_paths"],
        _decisions(),
        "ID",
        output_path,
        ctx,
    )

    blocked_payload = blocked.model_dump(mode="json", by_alias=True)
    assert blocked_payload["outcome"] == "needs_input"
    assert blocked_payload["next_action"] == "request_user_input"
    assert blocked_payload["retryable"] is False
    assert "prepared_input_relative_path" not in blocked_payload
    assert "input_identity" not in blocked_payload
    assert blocked_payload["requirements"] == [
        {
            "code": "required_fields_missing",
            "relative_path": "warehouses.csv",
            "unit_ref": "table",
            "candidate_roles": ["existing_warehouse"],
            "missing_required_fields": ["warehouse_type"],
            "field_name": None,
            "question": (
                "文件 warehouses.csv 的单元 table 缺少仓型字段 warehouse_type。"
                "请补充该列，每行使用 center 或 cross_docking，然后再继续。"
            ),
        }
    ]
    assert not (tmp_path / output_path).exists()


def test_needs_input_summary_and_requirements_are_bounded_with_total_count() -> None:
    requirements = [
        DataSourceRequirement(
            code="source_data_invalid",
            relative_path="source.csv",
            unit_ref=f"table-{index}",
            candidate_roles=[SourceRole.DEMAND],
            missing_required_fields=[],
            question=f"问题 {index}",
        )
        for index in range(70)
    ]
    payload = data_server._needs_input_result(requirements).model_dump(mode="json", by_alias=True)
    assert payload["requirement_count"] == 70
    assert payload["requirements_truncated"] is True
    assert len(payload["requirements"]) == 64
    assert len(payload["summary"]) < 500


def test_preparation_bounded_count_contract_rejects_underfilled_or_wrong_flags() -> None:
    requirement = {
        "code": "source_data_invalid",
        "relative_path": "source.csv",
        "unit_ref": "table",
        "candidate_roles": ["demand"],
        "missing_required_fields": [],
        "field_name": None,
        "question": "修正选中来源。",
    }
    needs_input = {
        "outcome": "needs_input",
        "summary": "发现问题。",
        "next_action": "request_user_input",
        "retryable": False,
        "requirements": [requirement],
        "requirement_count": 1,
        "requirements_truncated": False,
    }
    for count, items, truncated in ((2, [requirement], False), (10, [requirement] * 9, True), (1, [requirement], True)):
        with pytest.raises(ValidationError):
            DataPreparationNeedsInput.model_validate(
                {**needs_input, "requirements": items, "requirement_count": count, "requirements_truncated": truncated}
            )

    ready = {
        "outcome": "ready",
        "operation": "created",
        "summary": "完成。",
        "next_action": "handoff",
        "retryable": False,
        "prepared_input_relative_path": "outputs/warehouse-network/prepared/x.json",
        "input_identity": {"content_sha256": "0" * 64},
        "state": "ready",
        "role_counts": PreparationRoleCounts(
            demand=0,
            existing_warehouse=0,
            candidate_warehouse=0,
            current_assignment=0,
            route_quote=0,
            provided_route_fact=0,
        ).model_dump(mode="json"),
        "warnings": [],
        "warning_count": 0,
        "warnings_truncated": False,
        "candidate_warehouse_count": 0,
        "candidate_warehouses": [],
        "candidate_warehouses_truncated": False,
    }
    for field, value in (
        ("warning_count", 1),
        ("candidate_warehouse_count", 1),
    ):
        with pytest.raises(ValidationError):
            DataPreparationReady.model_validate({**ready, field: value})
    with pytest.raises(ValidationError):
        DataPreparationReady.model_validate({**ready, "warnings_truncated": True})


def test_complete_source_profile_uses_exact_total_without_repeated_preview_evidence(
    tmp_path, monkeypatch
) -> None:
    rows = [
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,longitude,latitude,"
        "is_existing"
    ]
    rows.extend(
        f"CAND-{index},Candidate {index},cross_docking,CITY-{index},City {index},"
        f"{106 + index / 10:.1f},{-6 - index / 10:.1f},false"
        for index in range(1, 13)
    )
    rows[7] = (
        "WH-CANDIDATE-BALIKPAPAN,Balikpapan Candidate Cross Docking,cross_docking,"
        "IDN-CITY-021,Balikpapan,116.996737,-1.202072,false"
    )
    (tmp_path / "candidate.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")
    _use_store(tmp_path, monkeypatch)

    result = data_server.inspect_workspace_sources(
        ["candidate.csv"],
        [SourceRole.CANDIDATE_WAREHOUSE],
        ctx=object(),
        country_code="ID",
    )
    assert result.structuredContent is not None
    inspection = result.structuredContent
    source = inspection["source_profile"]["sources"][0]
    assert "structure" not in source
    unit = source["units"][0]
    inspection_identity = SourceInspectionIdentity.model_validate(
        inspection["inspection_identity"]
    )
    inspected_paths = inspection["inspected_relative_paths"]

    assert unit["record_count"] == 12
    assert unit["record_count_exact"] is True
    assert "fields" not in unit
    assert "preview" not in unit
    assert "mapping_suggestions" not in unit
    assert "Balikpapan" not in json.dumps(inspection, ensure_ascii=False)
    resolved = unit["role_assessments"][0]
    assert resolved["state"] == "complete"
    assert resolved["resolved_mapping"]["warehouse_type"] == "warehouse_type"

    with pytest.raises(ProviderContractError, match="generated_output_path_invalid"):
        data_server.prepare_network_input(
            inspection_identity,
            inspected_paths,
            [
                    SourceSelection(
                        relative_path="candidate.csv",
                        unit_ref="table",
                        role=SourceRole.CANDIDATE_WAREHOUSE,
                )
            ],
            "ID",
            "candidate-catalog.json",
            _context(tmp_path),
        )

    prepared = data_server.prepare_network_input(
        inspection_identity,
        inspected_paths,
        [
            SourceSelection(
                relative_path="candidate.csv",
                unit_ref="table",
                role=SourceRole.CANDIDATE_WAREHOUSE,
            )
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/candidate-catalog.json",
        _context(tmp_path),
    )
    prepared_payload = prepared.model_dump(mode="json", by_alias=True)
    assert prepared_payload["outcome"] == "needs_input"
    assert any(item["code"] == "source_data_invalid" for item in prepared_payload["requirements"])


def test_inspection_derives_country_from_admin_metadata_without_guessing(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity\nC-1,Jakarta,10\n", encoding="utf-8"
    )
    (tmp_path / "warehouse.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,is_existing\n"
        "W-1,Jakarta,center,C-1,Jakarta,true\n",
        encoding="utf-8",
    )
    admin = {
        "country_code": "ID",
        "admin_level": "city",
        "schema_version": "administrative_catalog.v1",
        "rows": [],
    }
    (tmp_path / "admin-id.json").write_text(json.dumps(admin), encoding="utf-8")
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)

    derived = data_server.inspect_workspace_sources(
        ["demand.csv", "admin-id.json"],
        [SourceRole.DEMAND],
        ctx=ctx,
    ).structuredContent
    assert derived["country_code"] == "ID"
    admin_source = next(
        source for source in derived["source_profile"]["sources"] if source["relative_path"] == "admin-id.json"
    )
    assert admin_source["administrative_metadata"] == {
        key: admin[key] for key in ("country_code", "admin_level", "schema_version")
    }

    (tmp_path / "ordinary.json").write_text(
        json.dumps({"schema_version": "ordinary.v1", "rows": []}), encoding="utf-8"
    )
    mixed = data_server.inspect_workspace_sources(
        ["demand.csv", "ordinary.json", "admin-id.json"],
        [SourceRole.DEMAND],
        ctx=ctx,
    ).structuredContent
    assert mixed["country_code"] == "ID"

    missing = data_server.inspect_workspace_sources(
        ["demand.csv"], [SourceRole.DEMAND], ctx=ctx
    ).structuredContent
    assert missing["country_code"] is None
    assert "country_code_not_found" in missing["warnings"]

    conflict_admin = dict(admin, country_code="MY")
    (tmp_path / "admin-my.json").write_text(json.dumps(conflict_admin), encoding="utf-8")
    conflict = data_server.inspect_workspace_sources(
        ["demand.csv", "admin-id.json", "admin-my.json"],
        [SourceRole.DEMAND],
        ctx=ctx,
    ).structuredContent
    assert conflict["country_code"] is None
    assert "country_code_metadata_conflict" in conflict["warnings"]

    explicit = data_server.inspect_workspace_sources(
        ["demand.csv", "warehouse.csv", "admin-id.json"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="MY",
    ).structuredContent
    assert explicit["country_code"] == "MY"
    mismatch = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(explicit["inspection_identity"]),
        explicit["inspected_relative_paths"],
        [
            SourceSelection(relative_path="demand.csv", unit_ref="table", role=SourceRole.DEMAND),
            SourceSelection(
                relative_path="warehouse.csv", unit_ref="table", role=SourceRole.EXISTING_WAREHOUSE
            ),
        ],
        "MY",
        f"{PREPARED_OUTPUT_DIR}/admin-country-mismatch.json",
        ctx,
        administrative_catalog_relative_path="admin-id.json",
    ).model_dump(mode="json", by_alias=True)
    assert mismatch["outcome"] == "needs_input"
    assert mismatch["requirements"][0]["candidate_roles"] == ["administrative_catalog"]


def test_current_mock_sources_have_no_global_false_demand_blockers(tmp_path, monkeypatch) -> None:
    fixture = tmp_path / "fixture"
    fixture.mkdir()
    fixture_root = Path(__file__).parents[3] / "apps/web/scripts/fixtures/warehouse-network/mock_data"
    for path in fixture_root.iterdir():
        if path.suffix in {".csv", ".json"} and path.name != "manifest.json":
            (fixture / path.name).write_bytes(path.read_bytes())
    _use_store(fixture, monkeypatch)

    result = data_server.inspect_workspace_sources(
        sorted(path.name for path in fixture.iterdir() if path.suffix in {".csv", ".json"}),
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=object(),
        country_code="ID",
    )
    assert result.structuredContent is not None
    inspection = result.structuredContent
    assert inspection["outcome"] == "inspected"
    assessments = {
        source["relative_path"]: [
            assessment
            for unit in source["units"]
            for assessment in unit["role_assessments"]
        ]
        for source in inspection["source_profile"]["sources"]
    }
    assert all(
        not (
            assessment["role"] == "demand"
            and assessment["state"] == "partial"
        )
        for source_assessments in assessments.values()
        for assessment in source_assessments
    )
    assert assessments["existing-warehouses.csv"][0]["role"] == "existing_warehouse"
    assert assessments["candidate-warehouses.csv"][0]["role"] == "candidate_warehouse"


def test_model_visible_mock_inspection_is_compact_json_with_resolved_alias_mappings(
    tmp_path, monkeypatch
) -> None:
    fixture_root = Path(__file__).parents[3] / "apps/web/scripts/fixtures/warehouse-network/mock_data"
    required_paths = [
        "administrative-areas.json",
        "candidate-warehouses.csv",
        "demand-cities.csv",
        "existing-warehouses.csv",
        "route-quotes.csv",
    ]
    for path in fixture_root.iterdir():
        if path.name in {*required_paths, "population-snapshot.csv"}:
            (tmp_path / path.name).write_bytes(path.read_bytes())
    _use_store(tmp_path, monkeypatch)

    for paths in (required_paths, [*required_paths, "population-snapshot.csv"]):
        result = data_server.inspect_workspace_sources(
            paths,
            [
                SourceRole.DEMAND,
                SourceRole.EXISTING_WAREHOUSE,
                SourceRole.CANDIDATE_WAREHOUSE,
                SourceRole.ROUTE_QUOTE,
            ],
            ctx=object(),
            country_code="ID",
        )
        assert result.structuredContent is not None
        payload = result.structuredContent
        serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

        assert payload["outcome"] == "inspected"
        assert len(serialized) < data_server._workspace_intake.MAX_MODEL_INSPECTION_BYTES
        assert json.loads(serialized) == payload
        assert payload["inspection_identity"]["schemaVersion"] == "workspace_source_inspection.v2"
        assert payload["inspected_relative_paths"] == sorted(paths)
        assert payload["source_profile"]["schemaVersion"] == "workspace_source_profile.v3"
        assert payload["source_profile"]["source_count"] == len(paths)

        sources = {
            source["relative_path"]: source for source in payload["source_profile"]["sources"]
        }
        candidate_unit = sources["candidate-warehouses.csv"]["units"][0]
        candidate_assessment = candidate_unit["role_assessments"][0]
        assert candidate_assessment == {
            "role": "candidate_warehouse",
            "state": "complete",
            "confidence": 1.0,
            "ambiguous": False,
            "resolved_mapping": {
                "city_id": "city_id",
                "city_name": "city_name",
                "is_existing": "is_existing",
                "is_fixed": "is_fixed",
                "latitude": "latitude",
                "longitude": "longitude",
                "province_id": "province_id",
                "province_name": "province_name",
                "upstream_center_id": "upstream_center_id",
                "warehouse_id": "warehouse_id",
                "warehouse_name": "warehouse_name",
                "warehouse_type": "warehouse_type",
            },
        }
        assert candidate_unit["record_count"] == 12
        assert candidate_unit["record_count_exact"] is True
        assert "fields" not in candidate_unit
        assert "mapping_suggestions" not in candidate_unit

        admin_source = sources["administrative-areas.json"]
        assert admin_source["administrative_metadata"]["country_code"] == "ID"
        assert admin_source["units"][0]["record_count"] == 50
        if len(paths) == 6:
            population_unit = sources["population-snapshot.csv"]["units"][0]
            assert population_unit["role_assessments"] == []
            assert [field["name"] for field in population_unit["fields"]] == [
                "city_id",
                "city_name",
                "province_name",
                "population",
            ]
            assert all(len(field["representative_values"]) <= 1 for field in population_unit["fields"])


def test_current_mock_sources_prepare_from_full_units_and_keep_preview_tail_rows(
    tmp_path, monkeypatch
) -> None:
    fixture_root = Path(__file__).parents[3] / "apps/web/scripts/fixtures/warehouse-network/mock_data"
    selected_names = {
        "administrative-areas.json",
        "candidate-warehouses.csv",
        "demand-cities.csv",
        "existing-warehouses.csv",
        "route-quotes.csv",
    }
    for path in fixture_root.iterdir():
        if path.name in selected_names:
            (tmp_path / path.name).write_bytes(path.read_bytes())
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    inspected = data_server.inspect_workspace_sources(
        sorted(selected_names),
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
        country_code="ID",
    )
    assert inspected.structuredContent is not None
    profile = inspected.structuredContent
    selections = [
        SourceSelection(relative_path="demand-cities.csv", unit_ref="table", role=SourceRole.DEMAND),
        SourceSelection(
            relative_path="existing-warehouses.csv",
            unit_ref="table",
            role=SourceRole.EXISTING_WAREHOUSE,
        ),
        SourceSelection(
            relative_path="candidate-warehouses.csv",
            unit_ref="table",
            role=SourceRole.CANDIDATE_WAREHOUSE,
        ),
        SourceSelection(relative_path="route-quotes.csv", unit_ref="table", role=SourceRole.ROUTE_QUOTE),
    ]
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(profile["inspection_identity"]),
        profile["inspected_relative_paths"],
        selections,
        "ID",
        f"{PREPARED_OUTPUT_DIR}/mock-full.json",
        ctx,
        administrative_catalog_relative_path="administrative-areas.json",
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "ready"
    assert payload["role_counts"] == {
        "demand": 50,
        "existing_warehouse": 11,
        "candidate_warehouse": 12,
        "current_assignment": 0,
        "route_quote": 580,
        "provided_route_fact": 580,
    }
    prepared = json.loads(
        (tmp_path / payload["prepared_input_relative_path"]).read_text(encoding="utf-8")
    )
    assert len(prepared["demand_cities"]) == 50
    assert len(prepared["warehouses"]) == 23
    assert any(item["city_name"] == "Balikpapan" for item in prepared["warehouses"])


def test_fresh_prepared_input_is_reused_and_only_selected_raw_changes_invalidate_it(
    tmp_path, monkeypatch
) -> None:
    fixture_root = Path(__file__).parents[3] / "apps/web/scripts/fixtures/warehouse-network/mock_data"
    selected_names = {
        "administrative-areas.json",
        "candidate-warehouses.csv",
        "demand-cities.csv",
        "existing-warehouses.csv",
        "population-snapshot.csv",
        "route-quotes.csv",
    }
    for path in fixture_root.iterdir():
        if path.name in selected_names:
            (tmp_path / path.name).write_bytes(path.read_bytes())
    original_admin = (tmp_path / "administrative-areas.json").read_bytes()
    original_demand = (tmp_path / "demand-cities.csv").read_bytes()
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    raw_paths = sorted(selected_names)
    first = data_server.inspect_workspace_sources(
        raw_paths,
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    created = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(first["inspection_identity"]),
        first["inspected_relative_paths"],
        [
            SourceSelection(relative_path="demand-cities.csv", unit_ref="table", role=SourceRole.DEMAND),
            SourceSelection(relative_path="existing-warehouses.csv", unit_ref="table", role=SourceRole.EXISTING_WAREHOUSE),
            SourceSelection(relative_path="candidate-warehouses.csv", unit_ref="table", role=SourceRole.CANDIDATE_WAREHOUSE),
            SourceSelection(relative_path="route-quotes.csv", unit_ref="table", role=SourceRole.ROUTE_QUOTE),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/reuse.json",
        ctx,
        administrative_catalog_relative_path="administrative-areas.json",
    )
    created_payload = created.model_dump(mode="json", by_alias=True)
    assert created_payload["outcome"] == "ready"
    unselected = data_server.inspect_workspace_sources(
        raw_paths,
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert unselected["outcome"] == "inspected"
    reuse_paths = raw_paths + [created_payload["prepared_input_relative_path"]]
    reused = data_server.inspect_workspace_sources(
        reuse_paths,
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert reused["outcome"] == "prepared_ready"
    assert reused["operation"] == "reused"
    assert reused["prepared_input_relative_path"] == created_payload["prepared_input_relative_path"]
    assert "source_profile" not in reused
    reused_from_admin_country = data_server.inspect_workspace_sources(
        reuse_paths,
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
    ).structuredContent
    assert reused_from_admin_country["outcome"] == "prepared_ready"
    with pytest.raises(ProviderContractError, match="prepared_candidate_country_required"):
        data_server.inspect_workspace_sources(
            [created_payload["prepared_input_relative_path"]],
            [SourceRole.DEMAND],
            ctx=ctx,
        )

    canonical_copy = tmp_path / f"{PREPARED_OUTPUT_DIR}/a-copy.json"
    canonical_copy.write_bytes(
        (tmp_path / created_payload["prepared_input_relative_path"]).read_bytes()
    )
    canonical = data_server.inspect_workspace_sources(
        raw_paths + [
            created_payload["prepared_input_relative_path"],
            f"{PREPARED_OUTPUT_DIR}/a-copy.json",
        ],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert canonical["outcome"] == "prepared_ready"
    assert canonical["prepared_input_relative_path"] == f"{PREPARED_OUTPUT_DIR}/a-copy.json"
    canonical_copy.unlink()

    changed_payload = json.loads(
        (tmp_path / created_payload["prepared_input_relative_path"]).read_text(encoding="utf-8")
    )
    changed_payload["demand_cities"][0]["city_name"] = "Changed City"
    changed_path = tmp_path / f"{PREPARED_OUTPUT_DIR}/changed.json"
    changed_path.write_text(json.dumps(changed_payload), encoding="utf-8")
    ambiguous = data_server.inspect_workspace_sources(
        raw_paths + [
            created_payload["prepared_input_relative_path"],
            f"{PREPARED_OUTPUT_DIR}/changed.json",
        ],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert ambiguous["outcome"] == "prepared_selection_required"
    assert ambiguous["candidate_count"] == 2
    assert ambiguous["candidates_truncated"] is False
    assert "source_profile" not in ambiguous

    missing_role = data_server.inspect_workspace_sources(
        raw_paths + [created_payload["prepared_input_relative_path"]],
        [SourceRole.DEMAND, SourceRole.CURRENT_ASSIGNMENT],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert missing_role["outcome"] == "inspected"
    assert any("required_roles_missing" in warning for warning in missing_role["warnings"])
    wrong_country = data_server.inspect_workspace_sources(
        raw_paths + [created_payload["prepared_input_relative_path"]],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="MY",
    ).structuredContent
    assert wrong_country["outcome"] == "inspected"
    assert any("country_mismatch" in warning for warning in wrong_country["warnings"])
    needs_geography_payload = dict(changed_payload)
    needs_geography_payload["state"] = "needs_geography"
    needs_geography_path = tmp_path / f"{PREPARED_OUTPUT_DIR}/needs-geography.json"
    needs_geography_path.write_text(json.dumps(needs_geography_payload), encoding="utf-8")
    needs_geography = data_server.inspect_workspace_sources(
        raw_paths + [f"{PREPARED_OUTPUT_DIR}/needs-geography.json"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert needs_geography["outcome"] == "inspected"
    assert any("state_not_ready" in warning for warning in needs_geography["warnings"])

    mapping_payload = json.loads(
        (tmp_path / created_payload["prepared_input_relative_path"]).read_text(encoding="utf-8")
    )
    mapping_payload["source_selections"][0]["mappings"][0]["source_field"] = "missing_field"
    mapping_path = tmp_path / f"{PREPARED_OUTPUT_DIR}/mapping-mismatch.json"
    mapping_path.write_text(json.dumps(mapping_payload), encoding="utf-8")
    mapping_mismatch = data_server.inspect_workspace_sources(
        raw_paths + [f"{PREPARED_OUTPUT_DIR}/mapping-mismatch.json"],
        [SourceRole.DEMAND],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert mapping_mismatch["outcome"] == "inspected"
    assert any("prepared_selected_source_identity_invalid" in warning for warning in mapping_mismatch["warnings"])

    invalid_candidate_path = tmp_path / f"{PREPARED_OUTPUT_DIR}/invalid.json"
    invalid_candidate_path.write_text('{"schemaVersion":"prepared_network_input.v1"}', encoding="utf-8")
    invalid_candidate = data_server.inspect_workspace_sources(
        raw_paths + [f"{PREPARED_OUTPUT_DIR}/invalid.json"],
        [SourceRole.DEMAND],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert invalid_candidate["outcome"] == "inspected"
    assert any("prepared_network_input_invalid" in warning for warning in invalid_candidate["warnings"])

    with pytest.raises(ValueError, match="relative_paths must not contain duplicates"):
        data_server.inspect_workspace_sources(
            reuse_paths + [created_payload["prepared_input_relative_path"]],
            [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
            ctx=ctx,
            country_code="ID",
        )
    with pytest.raises(ValueError, match="prepared_candidate_not_found"):
        data_server.inspect_workspace_sources(
            [f"{PREPARED_OUTPUT_DIR}/missing.json"],
            [SourceRole.DEMAND],
            ctx=ctx,
            country_code="ID",
        )

    (tmp_path / "population-snapshot.csv").write_text(
        (tmp_path / "population-snapshot.csv").read_text() + "extra,Extra,Province,1\n",
        encoding="utf-8",
    )
    still_reused = data_server.inspect_workspace_sources(
        reuse_paths,
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert still_reused["outcome"] == "prepared_ready"

    (tmp_path / "administrative-areas.json").write_bytes(original_admin + b"\n")
    admin_stale = data_server.inspect_workspace_sources(
        reuse_paths,
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert admin_stale["outcome"] == "inspected"
    assert any("administrative_catalog_changed" in warning for warning in admin_stale["warnings"])
    (tmp_path / "administrative-areas.json").write_bytes(original_admin)

    (tmp_path / "demand-cities.csv").write_text(
        (tmp_path / "demand-cities.csv").read_text() + "\n",
        encoding="utf-8",
    )
    stale = data_server.inspect_workspace_sources(
        reuse_paths,
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    assert stale["outcome"] == "inspected"
    assert stale["warnings"]

    (tmp_path / "demand-cities.csv").unlink()
    with pytest.raises(ValueError, match="workspace_source_not_found"):
        data_server.inspect_workspace_sources(
            reuse_paths,
            [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
            ctx=ctx,
            country_code="ID",
        )
    (tmp_path / "demand-cities.csv").write_bytes(original_demand)
    outside = tmp_path / "outside-demand.csv"
    outside.write_bytes(original_demand)
    (tmp_path / "demand-cities.csv").unlink()
    (tmp_path / "demand-cities.csv").symlink_to(outside)
    with pytest.raises(ValueError, match="workspace_source_symlink_rejected"):
        data_server.inspect_workspace_sources(
            reuse_paths,
            [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE, SourceRole.CANDIDATE_WAREHOUSE, SourceRole.ROUTE_QUOTE],
            ctx=ctx,
            country_code="ID",
        )


def test_prepare_uses_one_post_read_snapshot_for_identity_and_provenance(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity,longitude,latitude\nC-1,Jakarta,10,106.8,-6.2\n",
        encoding="utf-8",
    )
    (tmp_path / "warehouse.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,longitude,latitude,is_existing\n"
        "W-1,Jakarta Center,center,C-1,Jakarta,106.8,-6.2,true\n",
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    inspection = data_server.inspect_workspace_sources(
        ["demand.csv", "warehouse.csv"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    original_snapshot = data_server.source_inspection_snapshot
    snapshot_calls = 0

    def changed_after_read(root, relative_paths):
        nonlocal snapshot_calls
        snapshot_calls += 1
        snapshot = original_snapshot(root, relative_paths)
        if snapshot_calls == 2:
            return SourceInspectionSnapshot(
                content_sha256="f" * 64,
                source_count=snapshot.source_count,
                file_sha256=snapshot.file_sha256,
            )
        return snapshot

    monkeypatch.setattr(data_server, "source_inspection_snapshot", changed_after_read)
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(inspection["inspection_identity"]),
        inspection["inspected_relative_paths"],
        [
            SourceSelection(relative_path="demand.csv", unit_ref="table", role=SourceRole.DEMAND),
            SourceSelection(
                relative_path="warehouse.csv", unit_ref="table", role=SourceRole.EXISTING_WAREHOUSE
            ),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/snapshot-race.json",
        ctx,
    )

    payload = result.model_dump(mode="json", by_alias=True)
    assert snapshot_calls == 2
    assert payload["outcome"] == "source_changed"
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/snapshot-race.json").exists()


def test_xlsx_and_json_units_are_not_aggregated(tmp_path, monkeypatch) -> None:
    workbook = Workbook()
    workbook.active.title = "需求"
    workbook.active.append(["city_id", "city_name", "demand_quantity"])
    workbook.active.append(["C-1", "Jakarta", 10])
    warehouses = workbook.create_sheet("仓库")
    warehouses.append(["warehouse_id", "warehouse_name", "warehouse_type", "city_id", "city_name", "is_existing"])
    warehouses.append(["W-1", "Jakarta", "center", "C-1", "Jakarta", True])
    workbook.save(tmp_path / "network.xlsx")
    (tmp_path / "nested.json").write_text(
        json.dumps(
            {
                "payload": {
                    "demandRows": [{"city_id": "C-1", "city_name": "Jakarta", "quantity": 10}],
                    "warehouses": [{"warehouse_id": "W-1", "warehouse_name": "Jakarta"}],
                    "nested": [{"tags": [{"name": "primary"}]}],
                }
            }
        ),
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    result = data_server.inspect_workspace_sources(
        ["network.xlsx", "nested.json"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=object(),
        country_code="ID",
    )
    assert result.structuredContent is not None
    sources = {item["relative_path"]: item for item in result.structuredContent["source_profile"]["sources"]}
    xlsx_units = {unit["unit_ref"]: unit for unit in sources["network.xlsx"]["units"]}
    assert set(xlsx_units) == {"sheet:需求", "sheet:仓库"}
    assert {unit["kind"] for unit in xlsx_units.values()} == {"sheet"}
    assert any(item["role"] == "demand" for item in xlsx_units["sheet:需求"]["role_assessments"])
    assert any(item["role"] == "existing_warehouse" for item in xlsx_units["sheet:仓库"]["role_assessments"])
    json_units = {unit["unit_ref"]: unit for unit in sources["nested.json"]["units"]}
    assert set(json_units) == {
        "$.payload.demandRows",
        "$.payload.warehouses",
        "$.payload.nested",
        "$.payload.nested[*].tags",
    }
    assert all(unit["kind"] == "json_array" for unit in json_units.values())
    assert json_units["$.payload.nested[*].tags"]["locator"] == {
        "path": "$.payload.nested[*].tags",
        "array_prefix": "payload.nested.item.tags",
    }
    xlsx_inspected = data_server.inspect_workspace_sources(
        ["network.xlsx"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=object(),
        country_code="ID",
    )
    assert xlsx_inspected.structuredContent is not None
    xlsx_profile = xlsx_inspected.structuredContent
    xlsx_result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(xlsx_profile["inspection_identity"]),
        xlsx_profile["inspected_relative_paths"],
        [
            SourceSelection(relative_path="network.xlsx", unit_ref="sheet:需求", role=SourceRole.DEMAND),
            SourceSelection(
                relative_path="network.xlsx", unit_ref="sheet:仓库", role=SourceRole.EXISTING_WAREHOUSE
            ),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/sheets.json",
        _context(tmp_path),
    )
    assert xlsx_result.model_dump(mode="json", by_alias=True)["outcome"] == "ready"


def test_explicit_mapping_handles_chinese_and_random_headers_without_name_inference(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "random.csv").write_text(
        "甲,乙,丙,丁,戊,己,庚\n"
        "C-1,Jakarta,10,-6.2,106.8,,\n",
        encoding="utf-8",
    )
    (tmp_path / "opaque.csv").write_text(
        "a1,a2,a3,a4,a5,a6\n"
        "WH-1,Jakarta Center,center,C-1,Jakarta,true\n",
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    inspected = data_server.inspect_workspace_sources(
        ["random.csv", "opaque.csv"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    )
    assert inspected.structuredContent is not None
    profile = inspected.structuredContent
    random_unit = profile["source_profile"]["sources"][0]["units"][0]
    assert random_unit["role_assessments"] == []
    assert [field["name"] for field in random_unit["fields"]] == [
        "甲",
        "乙",
        "丙",
        "丁",
        "戊",
        "己",
        "庚",
    ]
    assert random_unit["fields"][0] == {
        "name": "甲",
        "sample_types": ["string"],
        "representative_values": ["C-1"],
    }
    assert all(len(field["representative_values"]) <= 1 for field in random_unit["fields"])
    selections = [
        SourceSelection(
            relative_path="random.csv",
            unit_ref="table",
            role=SourceRole.DEMAND,
            mappings=[
                {"source_field": "甲", "target_field": "city_id", "transform": "trim"},
                {"source_field": "乙", "target_field": "city_name", "transform": "trim"},
                {"source_field": "丙", "target_field": "demand_quantity", "transform": "parse_integer"},
                {"source_field": "丁", "target_field": "latitude", "transform": "parse_decimal"},
                {"source_field": "戊", "target_field": "longitude", "transform": "parse_decimal"},
            ],
        ),
        SourceSelection(
            relative_path="opaque.csv",
            unit_ref="table",
            role=SourceRole.EXISTING_WAREHOUSE,
            mappings=[
                {"source_field": "a1", "target_field": "warehouse_id", "transform": "trim"},
                {"source_field": "a2", "target_field": "warehouse_name", "transform": "trim"},
                {"source_field": "a3", "target_field": "warehouse_type", "transform": "normalize_warehouse_type"},
                {"source_field": "a4", "target_field": "city_id", "transform": "trim"},
                {"source_field": "a5", "target_field": "city_name", "transform": "trim"},
                {"source_field": "a6", "target_field": "is_existing", "transform": "parse_boolean"},
            ],
        ),
    ]
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(profile["inspection_identity"]),
        profile["inspected_relative_paths"],
        selections,
        "ID",
        f"{PREPARED_OUTPUT_DIR}/opaque-mapped.json",
        ctx,
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "ready"
    assert payload["role_counts"]["demand"] == 1
    assert payload["role_counts"]["existing_warehouse"] == 1


def test_ambiguous_headers_keep_unit_local_field_and_mapping_evidence(tmp_path, monkeypatch) -> None:
    (tmp_path / "ambiguous.csv").write_text(
        "city_id,city_code,city_name,demand_quantity\nC-1,CITY-ONE,Jakarta,10\n",
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)

    result = data_server.inspect_workspace_sources(
        ["ambiguous.csv"],
        [SourceRole.DEMAND],
        ctx=object(),
        country_code="ID",
    )
    assert result.structuredContent is not None
    unit = result.structuredContent["source_profile"]["sources"][0]["units"][0]
    assessment = unit["role_assessments"][0]

    assert assessment["role"] == "demand"
    assert assessment["state"] == "ambiguous"
    assert assessment["confidence"] == 1.0
    assert assessment["ambiguous"] is True
    assert assessment["missing_required_fields"] == []
    city_id = next(
        mapping for mapping in assessment["candidate_mappings"] if mapping["target_field"] == "city_id"
    )
    assert city_id == {
        "target_field": "city_id",
        "source_fields": ["city_id", "city_code"],
        "transform": "normalize_identifier",
    }
    direct_mapping = SourceSelection.model_validate(
        {
            "relative_path": "ambiguous.csv",
            "unit_ref": "table",
            "role": "demand",
            "mappings": [
                {
                    "source_field": "city_id",
                    "target_field": "city_id",
                    "transform": city_id["transform"],
                },
                *[
                    {
                        "source_field": mapping["source_fields"][0],
                        "target_field": mapping["target_field"],
                        "transform": mapping["transform"],
                    }
                    for mapping in assessment["candidate_mappings"]
                    if mapping["target_field"] != "city_id"
                ],
            ],
        }
    )
    assert direct_mapping.mappings[0].transform.value == "normalize_identifier"
    assert [field["name"] for field in unit["fields"]] == [
        "city_id",
        "city_code",
        "city_name",
        "demand_quantity",
    ]
    assert "preview" not in unit


def test_one_csv_unit_can_supply_multiple_selected_roles_without_duplicate_reads(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "combined.csv").write_text(
        "city_id,city_name,demand_quantity,warehouse_id,warehouse_name,warehouse_type,is_existing\n"
        "C-1,Jakarta,10,W-1,Jakarta Center,center,true\n",
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    inspected = data_server.inspect_workspace_sources(
        ["combined.csv"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    )
    assert inspected.structuredContent is not None
    profile = inspected.structuredContent
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(profile["inspection_identity"]),
        profile["inspected_relative_paths"],
        [
            SourceSelection(
                relative_path="combined.csv",
                unit_ref="table",
                role=SourceRole.DEMAND,
                mappings=[
                    {"source_field": "city_id", "target_field": "city_id", "transform": "trim"},
                    {"source_field": "city_name", "target_field": "city_name", "transform": "trim"},
                    {"source_field": "demand_quantity", "target_field": "demand_quantity", "transform": "parse_integer"},
                ],
            ),
            SourceSelection(
                relative_path="combined.csv",
                unit_ref="table",
                role=SourceRole.EXISTING_WAREHOUSE,
                mappings=[
                    {"source_field": "warehouse_id", "target_field": "warehouse_id", "transform": "trim"},
                    {"source_field": "warehouse_name", "target_field": "warehouse_name", "transform": "trim"},
                    {"source_field": "warehouse_type", "target_field": "warehouse_type", "transform": "normalize_warehouse_type"},
                    {"source_field": "city_id", "target_field": "city_id", "transform": "trim"},
                    {"source_field": "city_name", "target_field": "city_name", "transform": "trim"},
                ],
            ),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/combined.json",
        ctx,
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "ready"
    assert payload["role_counts"]["demand"] == 1
    assert payload["role_counts"]["existing_warehouse"] == 1


def test_nested_json_array_unit_is_read_exactly_and_unselected_array_is_ignored(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "nested.json").write_text(
        json.dumps(
            {
                "payload": {
                    "demand": [{"x": "C-1", "y": "Jakarta", "q": 10}],
                    "unselected": [{"x": "BAD", "y": "Ignore", "q": "invalid"}],
                }
            }
        ),
        encoding="utf-8",
    )
    (tmp_path / "warehouse.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,is_existing\n"
        "WH-1,Jakarta Center,center,C-1,Jakarta,true\n",
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    inspected = data_server.inspect_workspace_sources(
        ["nested.json", "warehouse.csv"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    )
    assert inspected.structuredContent is not None
    profile = inspected.structuredContent
    selections = [
        SourceSelection(
            relative_path="nested.json",
            unit_ref="$.payload.demand",
            role=SourceRole.DEMAND,
            mappings=[
                {"source_field": "x", "target_field": "city_id", "transform": "trim"},
                {"source_field": "y", "target_field": "city_name", "transform": "trim"},
                {"source_field": "q", "target_field": "demand_quantity", "transform": "parse_integer"},
            ],
        ),
        SourceSelection(
            relative_path="warehouse.csv", unit_ref="table", role=SourceRole.EXISTING_WAREHOUSE
        ),
    ]
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(profile["inspection_identity"]),
        profile["inspected_relative_paths"],
        selections,
        "ID",
        f"{PREPARED_OUTPUT_DIR}/nested.json.input.json",
        ctx,
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "ready"
    prepared = json.loads((tmp_path / payload["prepared_input_relative_path"]).read_text())
    assert [item["city_id"] for item in prepared["demand_cities"]] == ["C-1"]


def test_selected_full_data_invalid_value_outside_preview_is_needs_input_no_write(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity\nC-1,Jakarta,10\n", encoding="utf-8"
    )
    rows = [
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,is_existing\n",
        "W-1,Jakarta,center,C-1,Jakarta,true\n",
        "W-2,Jakarta 2,center,C-1,Jakarta,true\n",
        "W-3,Jakarta 3,center,C-1,Jakarta,true\n",
        "W-4,Jakarta 4,not-a-type,C-1,Jakarta,true\n",
    ]
    (tmp_path / "warehouse.csv").write_text("".join(rows), encoding="utf-8")
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["demand.csv", "warehouse.csv"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    )
    assert profile.structuredContent is not None
    inspection = profile.structuredContent
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(inspection["inspection_identity"]),
        inspection["inspected_relative_paths"],
        [
            SourceSelection(relative_path="demand.csv", unit_ref="table", role=SourceRole.DEMAND),
            SourceSelection(
                relative_path="warehouse.csv", unit_ref="table", role=SourceRole.EXISTING_WAREHOUSE
            ),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/invalid-tail.json",
        ctx,
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "needs_input"
    assert any(item["code"] == "source_data_invalid" for item in payload["requirements"])
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/invalid-tail.json").exists()


def test_source_changed_covers_missing_selected_file_without_writing(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity\nC-1,Jakarta,10\n", encoding="utf-8"
    )
    (tmp_path / "warehouse.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name\n"
        "W-1,Jakarta,center,C-1,Jakarta\n",
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["demand.csv", "warehouse.csv"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    )
    assert profile.structuredContent is not None
    inspection = profile.structuredContent
    (tmp_path / "warehouse.csv").unlink()
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(inspection["inspection_identity"]),
        inspection["inspected_relative_paths"],
        [
            SourceSelection(relative_path="demand.csv", unit_ref="table", role=SourceRole.DEMAND),
            SourceSelection(
                relative_path="warehouse.csv", unit_ref="table", role=SourceRole.EXISTING_WAREHOUSE
            ),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/missing-source.json",
        ctx,
    )
    assert result.model_dump(mode="json", by_alias=True)["outcome"] == "source_changed"
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/missing-source.json").exists()


def test_invalid_source_unit_field_and_required_mapping_are_contract_errors(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity\nC-1,Jakarta,10\n", encoding="utf-8"
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["demand.csv"],
        [SourceRole.DEMAND],
        ctx=ctx,
        country_code="ID",
    )
    assert profile.structuredContent is not None
    inspection = profile.structuredContent
    identity = SourceInspectionIdentity.model_validate(inspection["inspection_identity"])
    paths = inspection["inspected_relative_paths"]
    base = dict(
        inspection_identity=identity,
        inspected_relative_paths=paths,
        country_code="ID",
        output_relative_path=f"{PREPARED_OUTPUT_DIR}/invalid.json",
        ctx=ctx,
    )
    with pytest.raises(ProviderContractError, match="source_unit_ref_invalid"):
        data_server.prepare_network_input(
            **base,
            source_selections=[
                SourceSelection(relative_path="demand.csv", unit_ref="$.missing", role=SourceRole.DEMAND)
            ],
        )
    with pytest.raises(ProviderContractError, match="source_selection_field_invalid"):
        data_server.prepare_network_input(
            **base,
            source_selections=[
                SourceSelection(
                    relative_path="demand.csv",
                    unit_ref="table",
                    role=SourceRole.DEMAND,
                    mappings=[
                        {"source_field": "not_a_field", "target_field": "city_id", "transform": "trim"},
                        {"source_field": "city_name", "target_field": "city_name", "transform": "trim"},
                        {"source_field": "demand_quantity", "target_field": "demand_quantity", "transform": "parse_integer"},
                    ],
                )
            ],
        )
    with pytest.raises(ProviderContractError, match="source_selection_required_mapping_missing"):
        data_server.prepare_network_input(
            **base,
            source_selections=[
                SourceSelection(
                    relative_path="demand.csv",
                    unit_ref="table",
                    role=SourceRole.DEMAND,
                    mappings=[
                        {"source_field": "city_id", "target_field": "city_id", "transform": "trim"},
                        {"source_field": "city_name", "target_field": "city_name", "transform": "trim"},
                    ],
                )
            ],
        )


def test_xlsx_selected_sheet_formula_is_a_full_data_blocker(tmp_path, monkeypatch) -> None:
    workbook = Workbook()
    demand = workbook.active
    demand.title = "demand"
    demand.append(["city_id", "city_name", "demand_quantity"])
    demand.append(["C-1", "Jakarta", 10])
    warehouse = workbook.create_sheet("warehouse")
    warehouse.append(["warehouse_id", "warehouse_name", "warehouse_type", "city_id", "city_name", "is_existing"])
    warehouse.append(["W-1", "Jakarta", "=\"center\"", "C-1", "Jakarta", True])
    workbook.save(tmp_path / "formula.xlsx")
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["formula.xlsx"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    )
    assert profile.structuredContent is not None
    inspection = profile.structuredContent
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(inspection["inspection_identity"]),
        inspection["inspected_relative_paths"],
        [
            SourceSelection(relative_path="formula.xlsx", unit_ref="sheet:demand", role=SourceRole.DEMAND),
            SourceSelection(
                relative_path="formula.xlsx", unit_ref="sheet:warehouse", role=SourceRole.EXISTING_WAREHOUSE
            ),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/formula.json",
        ctx,
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "needs_input"
    assert any(
        item.get("field_name") == "warehouse_type" for item in payload["requirements"]
    )
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/formula.json").exists()


def test_same_role_units_merge_exact_duplicates_and_block_conflicts_no_write(
    tmp_path, monkeypatch
) -> None:
    demand_header = "city_id,city_name,demand_quantity,longitude,latitude\n"
    demand_row = "C-1,Jakarta,10,106.8,-6.2\n"
    (tmp_path / "demand-a.csv").write_text(demand_header + demand_row, encoding="utf-8")
    (tmp_path / "demand-b.csv").write_text(demand_header + demand_row, encoding="utf-8")
    (tmp_path / "warehouse.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,is_existing,longitude,latitude\n"
        "W-1,Jakarta,center,C-1,Jakarta,true,106.8,-6.2\n",
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["demand-a.csv", "demand-b.csv", "warehouse.csv"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    selections = [
        SourceSelection(relative_path="demand-a.csv", unit_ref="table", role=SourceRole.DEMAND),
        SourceSelection(relative_path="demand-b.csv", unit_ref="table", role=SourceRole.DEMAND),
        SourceSelection(relative_path="warehouse.csv", unit_ref="table", role=SourceRole.EXISTING_WAREHOUSE),
    ]
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(profile["inspection_identity"]),
        profile["inspected_relative_paths"],
        selections,
        "ID",
        f"{PREPARED_OUTPUT_DIR}/duplicate.json",
        ctx,
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "ready"
    assert any("重复记录数：1" in warning for warning in payload["warnings"])

    (tmp_path / "demand-b.csv").write_text(
        demand_header + "C-1,Surabaya,10,106.8,-6.2\n", encoding="utf-8"
    )
    conflict_profile = data_server.inspect_workspace_sources(
        ["demand-a.csv", "demand-b.csv", "warehouse.csv"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    conflict = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(conflict_profile["inspection_identity"]),
        conflict_profile["inspected_relative_paths"],
        selections,
        "ID",
        f"{PREPARED_OUTPUT_DIR}/conflict.json",
        ctx,
    )
    conflict_payload = conflict.model_dump(mode="json", by_alias=True)
    assert conflict_payload["outcome"] == "needs_input"
    assert any(
        item["code"] == "source_duplicate_conflict"
        and item["relative_path"] == "demand-b.csv"
        for item in conflict_payload["requirements"]
    )
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/conflict.json").exists()


def test_geography_error_is_attributed_to_warehouse_unit_not_first_demand_unit(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity\nC-1,Jakarta,10\n", encoding="utf-8"
    )
    (tmp_path / "warehouse.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,is_existing\n"
        "W-1,Unknown,center,C-2,Unknown,true\n",
        encoding="utf-8",
    )
    (tmp_path / "admin.json").write_text(
        json.dumps(
            {
                "country_code": "ID",
                "admin_level": "city",
                "rows": [
                    {
                        "city_id": "C-1",
                        "city_name": "Jakarta",
                        "longitude": 106.8,
                        "latitude": -6.2,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["demand.csv", "warehouse.csv", "admin.json"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(profile["inspection_identity"]),
        profile["inspected_relative_paths"],
        [
            SourceSelection(relative_path="demand.csv", unit_ref="table", role=SourceRole.DEMAND),
            SourceSelection(relative_path="warehouse.csv", unit_ref="table", role=SourceRole.EXISTING_WAREHOUSE),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/geo-error.json",
        ctx,
        administrative_catalog_relative_path="admin.json",
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "needs_input"
    assert all(item["relative_path"] == "warehouse.csv" for item in payload["requirements"])
    assert all(item["candidate_roles"] == ["existing_warehouse"] for item in payload["requirements"])
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/geo-error.json").exists()


def test_candidate_geography_error_is_attributed_to_candidate_unit(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity\nC-1,Jakarta,10\n", encoding="utf-8"
    )
    (tmp_path / "candidate.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,is_existing\n"
        "W-CAND,Unknown,center,C-2,Unknown,false\n",
        encoding="utf-8",
    )
    (tmp_path / "admin.json").write_text(
        json.dumps(
            {
                "country_code": "ID",
                "admin_level": "city",
                "rows": [
                    {
                        "city_id": "C-1",
                        "city_name": "Jakarta",
                        "longitude": 106.8,
                        "latitude": -6.2,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["demand.csv", "candidate.csv", "admin.json"],
        [SourceRole.DEMAND, SourceRole.CANDIDATE_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(profile["inspection_identity"]),
        profile["inspected_relative_paths"],
        [
            SourceSelection(relative_path="demand.csv", unit_ref="table", role=SourceRole.DEMAND),
            SourceSelection(relative_path="candidate.csv", unit_ref="table", role=SourceRole.CANDIDATE_WAREHOUSE),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/candidate-geo-error.json",
        ctx,
        administrative_catalog_relative_path="admin.json",
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "needs_input"
    assert all(item["relative_path"] == "candidate.csv" for item in payload["requirements"])
    assert all(item["candidate_roles"] == ["candidate_warehouse"] for item in payload["requirements"])
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/candidate-geo-error.json").exists()


def test_administrative_catalog_exception_is_attributed_to_admin_document(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity\nC-1,Jakarta,10\n", encoding="utf-8"
    )
    (tmp_path / "warehouse.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,is_existing\n"
        "W-1,Jakarta,center,C-1,Jakarta,true\n",
        encoding="utf-8",
    )
    (tmp_path / "bad-admin.json").write_text(json.dumps({"rows": []}), encoding="utf-8")
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["demand.csv", "warehouse.csv", "bad-admin.json"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=ctx,
        country_code="ID",
    ).structuredContent
    result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(profile["inspection_identity"]),
        profile["inspected_relative_paths"],
        [
            SourceSelection(relative_path="demand.csv", unit_ref="table", role=SourceRole.DEMAND),
            SourceSelection(relative_path="warehouse.csv", unit_ref="table", role=SourceRole.EXISTING_WAREHOUSE),
        ],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/bad-admin.json",
        ctx,
        administrative_catalog_relative_path="bad-admin.json",
    )
    payload = result.model_dump(mode="json", by_alias=True)
    assert payload["outcome"] == "needs_input"
    assert payload["requirements"][0]["relative_path"] == "bad-admin.json"
    assert payload["requirements"][0]["candidate_roles"] == ["administrative_catalog"]
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/bad-admin.json").exists()


def test_inspection_limits_are_explicit_selection_errors(tmp_path, monkeypatch) -> None:
    for name in ("a.csv", "b.csv"):
        (tmp_path / name).write_text("city_id,city_name\nC-1,Jakarta\n", encoding="utf-8")
    _use_store(tmp_path, monkeypatch)
    with pytest.raises(ValueError, match="at least one Workspace-relative path"):
        data_server.inspect_workspace_sources(
            [],
            [SourceRole.DEMAND],
            ctx=object(),
            country_code="ID",
        )

    monkeypatch.setattr(data_server._workspace_intake, "MAX_INSPECTION_FILES", 1)
    file_limited = data_server.inspect_workspace_sources(
        ["a.csv", "b.csv"],
        [SourceRole.DEMAND],
        ctx=object(),
        country_code="ID",
    )
    assert file_limited.structuredContent is not None
    assert file_limited.structuredContent["outcome"] == "selection_required"
    assert set(file_limited.structuredContent) == {
        "outcome",
        "schemaVersion",
        "summary",
        "code",
        "next_action",
        "retryable",
        "observed",
        "limit",
    }
    assert file_limited.structuredContent["code"] == "inspection_selection_required"
    assert file_limited.structuredContent["observed"]["files"] == 2
    assert file_limited.structuredContent["limit"]["files"] == 1

    workbook = Workbook()
    workbook.active.append(["city_id", "city_name", "demand_quantity"])
    workbook.create_sheet("second").append(["warehouse_id"])
    workbook.save(tmp_path / "multi.xlsx")
    monkeypatch.setattr(data_server._workspace_intake, "MAX_INSPECTION_FILES", 64)
    monkeypatch.setattr(data_server._workspace_intake, "MAX_INSPECTION_UNITS", 1)
    unit_limited = data_server.inspect_workspace_sources(
        ["multi.xlsx"],
        [SourceRole.DEMAND],
        ctx=object(),
        country_code="ID",
    )
    assert unit_limited.structuredContent is not None
    assert unit_limited.structuredContent["outcome"] == "selection_required"
    assert unit_limited.structuredContent["observed"]["files"] == 1
    assert unit_limited.structuredContent["observed"]["units"] == 2
    assert unit_limited.structuredContent["limit"]["units"] == 1

    monkeypatch.setattr(data_server._workspace_intake, "MAX_INSPECTION_UNITS", 128)
    monkeypatch.setattr(data_server._workspace_intake, "MAX_MODEL_INSPECTION_BYTES", 1)
    byte_limited = data_server.inspect_workspace_sources(
        ["a.csv"],
        [SourceRole.DEMAND],
        ctx=object(),
        country_code="ID",
    )
    assert byte_limited.structuredContent is not None
    assert byte_limited.structuredContent["outcome"] == "selection_required"
    assert byte_limited.structuredContent["observed"]["bytes"] > 1
    assert byte_limited.structuredContent["limit"]["bytes"] == 1


def test_prepare_input_writes_a_complete_auditable_workspace_document(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "demand.csv").write_text(
        "city_id,city_name,demand_quantity\nCITY-1,Jakarta,50\n",
        encoding="utf-8",
    )
    (tmp_path / "warehouses.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,is_existing\n"
        "WH-1,Jakarta Center,center,CITY-1,Jakarta,true\n",
        encoding="utf-8",
    )
    (tmp_path / "admin.json").write_text(
        json.dumps(
            {
                "country_code": "ID",
                "admin_level": "city",
                "rows": [
                    {
                        "city_id": "CITY-1",
                        "city_name": "Jakarta",
                        "province_id": "PROV-1",
                        "province_name": "Jakarta",
                        "longitude": 106.8,
                        "latitude": -6.2,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["demand.csv", "warehouses.csv", "admin.json"],
        [SourceRole.DEMAND, SourceRole.EXISTING_WAREHOUSE],
        ctx=object(),
        country_code="ID",
    )
    assert profile.structuredContent is not None
    inspection = profile.structuredContent
    inspection_identity = SourceInspectionIdentity.model_validate(
        inspection["inspection_identity"]
    )
    inspected_paths = inspection["inspected_relative_paths"]

    with pytest.raises(ValueError, match="country_code_required_iso_alpha2"):
        data_server.prepare_network_input(
            inspection_identity,
            inspected_paths,
            _decisions(),
            "IDN",
            f"{PREPARED_OUTPUT_DIR}/prepared-input.json",
            ctx,
        )

    result = data_server.prepare_network_input(
        inspection_identity,
        inspected_paths,
        _decisions(),
        "ID",
        f"{PREPARED_OUTPUT_DIR}/prepared-input.json",
        ctx,
    )
    assert isinstance(result, DataPreparationToolResult)
    result_payload = result.model_dump(mode="json", by_alias=True)
    assert result_payload["outcome"] == "ready"
    assert result_payload["state"] == "needs_geography"
    assert result_payload["prepared_input_relative_path"] == f"{PREPARED_OUTPUT_DIR}/prepared-input.json"
    payload = json.loads((tmp_path / result_payload["prepared_input_relative_path"]).read_text())
    assert payload["schemaVersion"] == "prepared_network_input.v2"
    assert all(source["mappings"] for source in payload["source_selections"])
    assert payload["issues"]
    assert result_payload["input_identity"]["content_sha256"]

    original_admin = (tmp_path / "admin.json").read_bytes()
    original_admin_reader = data_server.read_json_document_with_sha256

    def mutate_admin_before_enrichment(root, relative_path):
        if relative_path == "admin.json":
            changed = json.loads((root / relative_path).read_text(encoding="utf-8"))
            changed["rows"][0]["longitude"] = 107.1
            (root / relative_path).write_text(json.dumps(changed), encoding="utf-8")
        return original_admin_reader(root, relative_path)

    monkeypatch.setattr(data_server, "read_json_document_with_sha256", mutate_admin_before_enrichment)
    raced_primary = data_server.prepare_network_input(
        inspection_identity,
        inspected_paths,
        _decisions(),
        "ID",
        f"{PREPARED_OUTPUT_DIR}/prepared-input-admin-race.json",
        ctx,
        administrative_catalog_relative_path="admin.json",
    )
    raced_primary_payload = raced_primary.model_dump(mode="json", by_alias=True)
    assert raced_primary_payload["outcome"] == "source_changed"
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/prepared-input-admin-race.json").exists()
    (tmp_path / "admin.json").write_bytes(original_admin)
    monkeypatch.setattr(data_server, "read_json_document_with_sha256", original_admin_reader)

    original_demand = (tmp_path / "demand.csv").read_bytes()
    (tmp_path / "demand.csv").write_bytes(original_demand + b"\n")
    stale_geography = data_server.prepare_network_geography(
        result_payload["prepared_input_relative_path"],
        "admin.json",
        f"{PREPARED_OUTPUT_DIR}/prepared-input-geography-stale.json",
        ctx,
    )
    assert stale_geography.model_dump(mode="json", by_alias=True)["outcome"] == "source_changed"
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/prepared-input-geography-stale.json").exists()
    (tmp_path / "demand.csv").write_bytes(original_demand)

    updated_admin = json.loads((tmp_path / "admin.json").read_text(encoding="utf-8"))
    updated_admin["rows"][0]["longitude"] = 107.0
    (tmp_path / "admin-updated.json").write_text(
        json.dumps(updated_admin), encoding="utf-8"
    )
    updated_geography = data_server.prepare_network_geography(
        result_payload["prepared_input_relative_path"],
        "admin-updated.json",
        f"{PREPARED_OUTPUT_DIR}/prepared-input-geography-updated.json",
        ctx,
    )
    updated_payload = updated_geography.model_dump(mode="json", by_alias=True)
    assert updated_payload["outcome"] == "ready"
    updated_document = json.loads(
        (tmp_path / updated_payload["prepared_input_relative_path"]).read_text(encoding="utf-8")
    )
    assert updated_document["administrative_catalog"]["relative_path"] == "admin-updated.json"
    assert updated_document["demand_cities"][0]["longitude"] == 107.0

    catalog_reads: list[str] = []

    def counted_admin_reader(root, relative_path):
        catalog_reads.append(relative_path)
        return original_admin_reader(root, relative_path)

    monkeypatch.setattr(data_server, "read_json_document_with_sha256", counted_admin_reader)
    geography = data_server.prepare_network_geography(
        result_payload["prepared_input_relative_path"],
        "admin.json",
        f"{PREPARED_OUTPUT_DIR}/prepared-input-geography.json",
        ctx,
    )
    monkeypatch.setattr(data_server, "read_json_document_with_sha256", original_admin_reader)
    geography_payload = geography.model_dump(mode="json", by_alias=True)
    assert geography_payload["state"] == "ready"
    enriched = json.loads((tmp_path / geography_payload["prepared_input_relative_path"]).read_text())
    assert enriched["demand_cities"][0]["longitude"] == 106.8
    assert enriched["selected_source_identity"]
    assert geography_payload["input_identity"] != result_payload["input_identity"]
    assert catalog_reads == ["admin.json"]
    assert enriched["administrative_catalog"]["content_sha256"] == hashlib.sha256(
        original_admin
    ).hexdigest()

    atomic = data_server.prepare_network_input(
        inspection_identity,
        inspected_paths,
        _decisions(),
        "ID",
        f"{PREPARED_OUTPUT_DIR}/prepared-input-atomic.json",
        ctx,
        administrative_catalog_relative_path="admin.json",
    )
    atomic_payload_result = atomic.model_dump(mode="json", by_alias=True)
    assert atomic_payload_result["state"] == "ready"
    atomic_payload = json.loads((tmp_path / atomic_payload_result["prepared_input_relative_path"]).read_text())
    assert atomic_payload["demand_cities"][0]["longitude"] == 106.8
    assert atomic_payload["administrative_catalog"]["relative_path"] == "admin.json"

    manual_demand = SourceSelection.model_validate(
        {
            "relative_path": "demand.csv",
            "unit_ref": "table",
            "role": "demand",
            "mappings": [
                {
                    "source_field": field,
                    "target_field": field,
                    "transform": transform,
                }
                for field, transform in (
                    ("city_id", "normalize_identifier"),
                    ("city_name", "trim"),
                    ("demand_quantity", "parse_integer"),
                )
            ],
        }
    )
    explicit_result = data_server.prepare_network_input(
        inspection_identity,
        inspected_paths,
        [manual_demand, _decisions()[1]],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/prepared-input-manual.json",
        ctx,
        administrative_catalog_relative_path="admin.json",
    )
    assert explicit_result.model_dump(mode="json", by_alias=True)["outcome"] == "ready"

    with pytest.raises(ProviderContractError, match="workspace_file_invalid"):
        data_server.prepare_network_input(
            inspection_identity,
            inspected_paths,
            _decisions(),
            "ID",
            f"{PREPARED_OUTPUT_DIR}/prepared-input.json",
            ctx,
        )


def test_candidate_changes_require_a_complete_new_workspace_preparation(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "candidate.csv").write_text(
        "warehouse_id,warehouse_name,warehouse_type,city_id,city_name,longitude,latitude,"
        "is_existing\n"
        "CAND-1,Jakarta Candidate,center,CITY-1,Jakarta,106.9,-6.3,false\n",
        encoding="utf-8",
    )
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)
    profile = data_server.inspect_workspace_sources(
        ["candidate.csv"],
        [SourceRole.CANDIDATE_WAREHOUSE],
        ctx=object(),
        country_code="ID",
    )
    assert profile.structuredContent is not None
    inspection = profile.structuredContent
    inspection_identity = SourceInspectionIdentity.model_validate(
        inspection["inspection_identity"]
    )
    inspected_paths = inspection["inspected_relative_paths"]
    candidate = SourceSelection(
        relative_path="candidate.csv",
        unit_ref="table",
        role=SourceRole.CANDIDATE_WAREHOUSE,
    )

    result = data_server.prepare_network_input(
        inspection_identity,
        inspected_paths,
        [candidate],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/candidate-only-preparation.json",
        ctx,
    )
    result_payload = result.model_dump(mode="json", by_alias=True)
    assert result_payload["outcome"] == "needs_input"
    assert "prepared_input_relative_path" not in result_payload
    assert not hasattr(data_server, "normalize_candidate_delta")
    assert not hasattr(data_server, "derive_normalized_network_input")


def test_inspection_identity_is_order_independent_and_rejects_changed_inputs(
    tmp_path, monkeypatch
) -> None:
    (tmp_path / "a.csv").write_text("city_id,city_name\nCITY-1,Jakarta\n", encoding="utf-8")
    (tmp_path / "b.csv").write_text("city_id,city_name\nCITY-2,Balikpapan\n", encoding="utf-8")
    _use_store(tmp_path, monkeypatch)
    ctx = _context(tmp_path)

    first = data_server.inspect_workspace_sources(
        ["b.csv", "a.csv"],
        [SourceRole.DEMAND],
        ctx=ctx,
        country_code="ID",
    )
    second = data_server.inspect_workspace_sources(
        ["a.csv", "b.csv"],
        [SourceRole.DEMAND],
        ctx=ctx,
        country_code="ID",
    )
    first_identity = first.structuredContent["inspection_identity"]
    assert first_identity == second.structuredContent["inspection_identity"]
    assert first.structuredContent["inspected_relative_paths"] == ["a.csv", "b.csv"]

    (tmp_path / "b.csv").write_text(
        "city_id,city_name\nCITY-2,Balikpapan changed\n", encoding="utf-8"
    )
    byte_result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(first_identity),
        ["a.csv", "b.csv"],
        [SourceSelection(relative_path="a.csv", unit_ref="table", role=SourceRole.DEMAND)],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/must-not-write.json",
        ctx,
    )
    assert byte_result.model_dump(mode="json", by_alias=True)["outcome"] == "source_changed"
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/must-not-write.json").exists()

    (tmp_path / "c.csv").write_text(
        "city_id,city_name\nCITY-3,Surabaya\n", encoding="utf-8"
    )
    path_set_result = data_server.prepare_network_input(
        SourceInspectionIdentity.model_validate(first_identity),
        ["a.csv", "b.csv", "c.csv"],
        [SourceSelection(relative_path="a.csv", unit_ref="table", role=SourceRole.DEMAND)],
        "ID",
        f"{PREPARED_OUTPUT_DIR}/must-not-write-path-set.json",
        ctx,
    )
    assert path_set_result.model_dump(mode="json", by_alias=True)["outcome"] == "source_changed"
    assert not (tmp_path / f"{PREPARED_OUTPUT_DIR}/must-not-write-path-set.json").exists()


def test_inspection_rejects_symlink_and_data_resource_surface_is_empty(tmp_path, monkeypatch) -> None:
    (tmp_path / "source.csv").write_text("city_id,city_name\nCITY-1,Jakarta\n", encoding="utf-8")
    (tmp_path / "outside.csv").write_text("city_id,city_name\nCITY-2,Outside\n", encoding="utf-8")
    (tmp_path / "linked.csv").symlink_to(tmp_path / "outside.csv")
    _use_store(tmp_path, monkeypatch)
    with pytest.raises(ValueError, match="workspace_source_symlink_rejected"):
        data_server.inspect_workspace_sources(
            ["linked.csv"],
            [SourceRole.DEMAND],
            ctx=_context(tmp_path),
            country_code="ID",
        )
    with pytest.raises(ValueError, match="workspace_source_symlink_rejected"):
        data_server.inspect_workspace_sources(
            ["../outside.csv"],
            [SourceRole.DEMAND],
            ctx=_context(tmp_path),
            country_code="ID",
        )
    assert asyncio.run(data_server.mcp.list_resources()) == []
    assert asyncio.run(data_server.mcp.list_resource_templates()) == []
