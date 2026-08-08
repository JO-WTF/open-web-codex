"""Bounded Workspace source discovery and evidence profiling.

This module deliberately exposes only structural samples to the model.  Full
file reads are reserved for the deterministic normalizer after the user has
confirmed the mapping after the requirement profile was published.
"""

from __future__ import annotations

import csv
import hashlib
import io
import itertools
import json
import re
import uuid
import zipfile
from pathlib import Path
from typing import Any, BinaryIO
from urllib.parse import unquote, urlparse
from xml.etree import ElementTree

import ijson
from charset_normalizer import from_bytes
from openpyxl import load_workbook

MAX_FILES = 500
MAX_BYTES = 100 * 1024 * 1024
MAX_SAMPLE_ROWS = 20
MAX_JSON_DEPTH = 64
MAX_JSON_NODES = 250_000
MAX_JSON_STRING = 64 * 1024
MAX_XLSX_SHEETS = 32
MAX_XLSX_COLUMNS = 256
MAX_NORMALIZE_ROWS = 1_000_000
EXCLUDED_DIRS = {
    ".git",
    "node_modules",
    "target",
    "dist",
    "build",
    ".venv",
    ".demo-staging",
    "__pycache__",
    "datasets",
}
SUPPORTED_SUFFIXES = {".xlsx", ".csv", ".json"}
SANDBOX_META = "codex/sandbox-state-meta"
DEMO_MANIFEST_SCHEMA = "demo_workspace_sources.v1"


def trusted_workspace_root(meta: Any) -> Path:
    extra = getattr(meta, "model_extra", None)
    state = extra.get(SANDBOX_META) if isinstance(extra, dict) else None
    sandbox_cwd = state.get("sandboxCwd") if isinstance(state, dict) else None
    if not isinstance(sandbox_cwd, str):
        raise ValueError("trusted Turn Workspace metadata is unavailable")
    parsed = urlparse(sandbox_cwd)
    if parsed.scheme != "file" or parsed.netloc not in ("", "localhost"):
        raise ValueError("trusted Turn Workspace is not a local file URI")
    root = Path(unquote(parsed.path)).resolve(strict=True)
    if not root.is_dir():
        raise ValueError("trusted Turn Workspace is not a directory")
    return root


def _iter_files(root: Path) -> list[Path]:
    files: list[Path] = []
    for path in root.rglob("*"):
        relative = path.relative_to(root)
        if any(part in EXCLUDED_DIRS for part in relative.parts):
            continue
        if path.is_symlink():
            continue
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            files.append(path)
    files.sort(key=lambda item: item.relative_to(root).as_posix())
    return files


def workspace_contains_supported_sources(root: Path) -> bool:
    """Detect any supported source, including files discovery would reject by size."""
    for path in root.rglob("*"):
        relative = path.relative_to(root)
        if any(part in EXCLUDED_DIRS for part in relative.parts):
            continue
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        if path.is_symlink():
            raise ValueError("workspace_supported_source_symlink_rejected")
        if path.is_file():
            return True
    return False


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def discover(root: Path) -> list[dict[str, Any]]:
    legacy = [
        path
        for path in root.rglob("*")
        if not path.is_symlink()
        and not any(part in EXCLUDED_DIRS for part in path.relative_to(root).parts)
        and path.is_file()
        and path.suffix.lower() in {".xls", ".xlsm"}
    ]
    if legacy:
        raise ValueError(
            "unsupported_source_format: legacy .xls/.xlsm files must be exported "
            "as .xlsx, .csv or .json"
        )
    files = _iter_files(root)
    records: list[dict[str, Any]] = []
    for path in files[:MAX_FILES]:
        size = path.stat().st_size
        if size <= 0 or size > MAX_BYTES:
            continue
        digest = _sha256_file(path)
        relative = path.relative_to(root).as_posix()
        source_ref = "source-" + hashlib.sha256(f"{relative}:{size}:{digest}".encode()).hexdigest()
        record = {
            "source_ref": source_ref,
            "display_name": path.name,
            "media_type": media_type(path),
            "byte_size": size,
            "content_sha256": digest,
            "extension": path.suffix.lower(),
        }
        parts = relative.split("/")
        if len(parts) >= 4 and parts[0:2] == [".open-web-codex", "source-assets"]:
            try:
                uuid.UUID(parts[2])
            except ValueError:
                pass
            else:
                record["source_asset_id"] = parts[2]
        records.append(record)
    if len(files) > MAX_FILES:
        raise ValueError("workspace_source_limit_exceeded: more than 500 supported files")
    return records


def workspace_source_metadata(root: Path) -> dict[str, Any]:
    """Return verified source classification without trusting a filename alone."""
    manifests = sorted(root.glob("demo-data/*/demo-manifest.json"))
    if not manifests:
        return {"dataClassification": "workspace_data"}
    if len(manifests) != 1 or manifests[0].is_symlink():
        raise ValueError("demo_manifest_conflict")
    manifest_path = manifests[0]
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError("demo_manifest_invalid") from error
    if (
        manifest.get("schemaVersion") != DEMO_MANIFEST_SCHEMA
        or manifest.get("dataClassification") != "synthetic_demo"
        or not isinstance(manifest.get("template"), dict)
        or not isinstance(manifest.get("files"), list)
    ):
        raise ValueError("demo_manifest_invalid")
    expected_names: set[str] = set()
    for item in manifest["files"]:
        if not isinstance(item, dict):
            raise ValueError("demo_manifest_invalid")
        name = item.get("displayName")
        expected_size = item.get("bytes")
        expected_digest = item.get("contentSha256")
        if (
            not isinstance(name, str)
            or Path(name).name != name
            or name in expected_names
            or not isinstance(expected_size, int)
            or not isinstance(expected_digest, str)
            or not re.fullmatch(r"[0-9a-f]{64}", expected_digest)
        ):
            raise ValueError("demo_manifest_invalid")
        expected_names.add(name)
        source = manifest_path.parent / name
        if source.is_symlink() or not source.is_file():
            raise ValueError("demo_manifest_source_missing")
        if source.stat().st_size != expected_size or _sha256_file(source) != expected_digest:
            raise ValueError("demo_manifest_source_mismatch")
    actual_names = {
        path.name
        for path in manifest_path.parent.iterdir()
        if path.is_file() and path.name != manifest_path.name
    }
    if actual_names != expected_names:
        raise ValueError("demo_manifest_source_mismatch")
    template = manifest["template"]
    return {
        "dataClassification": "synthetic_demo",
        "demoTemplate": {
            "id": template.get("id"),
            "version": template.get("version"),
            "seed": template.get("seed"),
        },
    }


def _resolve_ref(root: Path, source_ref: str) -> Path:
    for path in _iter_files(root):
        size = path.stat().st_size
        if size <= 0 or size > MAX_BYTES:
            continue
        digest = _sha256_file(path)
        relative = path.relative_to(root).as_posix()
        candidate = "source-" + hashlib.sha256(f"{relative}:{size}:{digest}".encode()).hexdigest()
        if candidate == source_ref:
            return path
    raise ValueError("source_ref is not present in the authorized Workspace")


def media_type(path: Path) -> str:
    return {
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".csv": "text/csv",
        ".json": "application/json",
    }[path.suffix.lower()]


def inspect(root: Path, source_ref: str) -> dict[str, Any]:
    path = _resolve_ref(root, source_ref)
    record = next(item for item in discover(root) if item["source_ref"] == source_ref)
    if path.suffix.lower() == ".csv":
        structure = inspect_csv(path)
    elif path.suffix.lower() == ".json":
        structure = inspect_json(path)
    else:
        structure = inspect_xlsx(path)
    return {**record, "structure": structure}


def inspect_csv(path: Path) -> dict[str, Any]:
    raw = _read_prefix(path, 256 * 1024)
    text = decode_text(raw)
    sample = text.splitlines()[: MAX_SAMPLE_ROWS + 1]
    if not sample:
        return {
            "kind": "table",
            "columns": [],
            "preview": _preview_payload([]),
            "record_count": 0,
            "record_count_exact": True,
        }
    try:
        dialect = csv.Sniffer().sniff("\n".join(sample[:10]), delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    rows = list(csv.reader(sample, delimiter=delimiter))
    header = [cell.strip() for cell in rows[0][:MAX_XLSX_COLUMNS]]
    preview_rows = [row[:MAX_XLSX_COLUMNS] for row in rows[1 : MAX_SAMPLE_ROWS + 1]]
    return {
        "kind": "table",
        "delimiter": delimiter,
        "columns": header,
        "preview": _preview_payload(preview_rows),
        "record_count": _count_csv_records(path, delimiter),
        "record_count_exact": True,
    }


def _preview_payload(rows: list[Any]) -> dict[str, Any]:
    return {
        "strategy": "head",
        "limit": MAX_SAMPLE_ROWS,
        "returned_count": len(rows),
        "complete": False,
        "rows": rows,
    }


def _count_csv_records(path: Path, delimiter: str) -> int:
    with path.open("rb") as binary:
        text_stream = io.TextIOWrapper(
            binary,
            encoding=detect_encoding(_read_prefix(path, 64 * 1024)),
            errors="strict",
        )
        reader = csv.reader(text_stream, delimiter=delimiter)
        next(reader, None)
        return sum(1 for row in reader if any(cell not in (None, "") for cell in row))


def inspect_json(path: Path) -> dict[str, Any]:
    arrays: dict[str, dict[str, Any]] = {}
    object_keys: dict[str, set[str]] = {}
    node_count = 0
    root_kind = "unknown"
    with path.open("rb") as stream:
        for prefix, event, value in _iter_json_events(stream):
            node_count += 1
            if node_count > MAX_JSON_NODES:
                raise ValueError("json_node_limit_exceeded")
            depth = _json_depth(prefix)
            if depth > MAX_JSON_DEPTH:
                raise ValueError("json_depth_limit_exceeded")
            if event == "start_map" and not prefix:
                root_kind = "object"
            elif event == "start_array" and not prefix:
                root_kind = "array"
            if event == "string" and isinstance(value, str) and len(value) > MAX_JSON_STRING:
                raise ValueError("json_string_limit_exceeded")
            if event == "start_array":
                array_key = prefix or "$"
                arrays[array_key] = {
                    "path": "$" if array_key == "$" else f"$.{array_key.replace('.item', '')}",
                    "length": 0,
                    "length_exact": True,
                    "preview": _preview_payload([]),
                }
            if event == "map_key":
                object_keys.setdefault(prefix, set()).add(str(value)[:256])

    # ijson's event stream is used above for limits and structure.  A second
    # streaming pass counts every item while extracting at most three preview
    # items per discovered array.  The full JSON document is never loaded into
    # memory or sent to the model.
    for array_prefix, array in arrays.items():
        item_prefix = "item" if array_prefix == "$" else f"{array_prefix}.item"
        with path.open("rb") as stream:
            try:
                for index, item in enumerate(ijson.items(stream, item_prefix)):
                    array["length"] = index + 1
                    if index < 3:
                        if isinstance(item, dict):
                            array["preview"]["rows"].append(
                                {
                                    "kind": "object",
                                    "fields": {
                                        str(key)[:256]: {
                                            "type": type(value).__name__,
                                            "sample": _bounded_json_sample(value),
                                        }
                                        for key, value in list(item.items())[:MAX_XLSX_COLUMNS]
                                    },
                                }
                            )
                        else:
                            array["preview"]["rows"].append(
                                {
                                    "kind": "value",
                                    "type": type(item).__name__,
                                    "sample": _bounded_json_sample(item),
                                }
                            )
                    array["preview"]["returned_count"] = len(array["preview"]["rows"])
            except (ijson.common.IncompleteJSONError, ijson.common.JSONError) as error:
                raise ValueError("invalid_json") from error
    return {
        "kind": "json",
        "tree": {"kind": root_kind},
        "object_keys": {
            path: sorted(keys)[:MAX_XLSX_COLUMNS] for path, keys in list(object_keys.items())[:256]
        },
        "arrays": list(arrays.values())[:256],
        "node_count": node_count,
    }


def _bounded_json_sample(value: Any) -> Any:
    if isinstance(value, str):
        return value[:256]
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    if isinstance(value, list):
        return [_bounded_json_sample(item) for item in value[:3]]
    if isinstance(value, dict):
        return {
            str(key)[:128]: _bounded_json_sample(item) for key, item in list(value.items())[:16]
        }
    return str(value)[:256]


def inspect_xlsx(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 2048:
            raise ValueError("xlsx_entry_limit_exceeded")
        expanded = sum(info.file_size for info in infos)
        if expanded > 512 * 1024 * 1024:
            raise ValueError("xlsx_expanded_size_limit_exceeded")
        if any(
            info.file_size and info.compress_size and info.file_size / info.compress_size > 100
            for info in infos
        ):
            raise ValueError("xlsx_compression_ratio_limit_exceeded")
        names = {info.filename for info in infos}
        if "xl/vbaProject.bin" in names or any(
            name.startswith("xl/externalLinks/") or name.endswith("externalLink1.xml")
            for name in names
        ):
            raise ValueError("unsupported_source_security_feature")
        _xlsx_shared_strings(archive)
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        # Do not echo the Runtime Workspace path into an Agent message.
        raise ValueError(f"xlsx_parse_failed:{type(error).__name__}") from error
    summaries = []
    try:
        for worksheet in list(workbook.worksheets)[:MAX_XLSX_SHEETS]:
            header: list[Any] | None = None
            preview_rows: list[list[Any]] = []
            record_count = 0
            for row in worksheet.iter_rows(values_only=False):
                values = []
                for cell in row[:MAX_XLSX_COLUMNS]:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        value = {"formula": True, "display": value[:MAX_JSON_STRING]}
                    values.append(value)
                if any(value not in (None, "") for value in values):
                    if header is None:
                        header = values
                    else:
                        record_count += 1
                        if len(preview_rows) < MAX_SAMPLE_ROWS:
                            preview_rows.append(values)
            if header is not None:
                summaries.append(
                    {
                        "sheet": worksheet.title,
                        "columns": [str(value or "").strip() for value in header],
                        "preview": _preview_payload(preview_rows),
                        "record_count": record_count,
                        "record_count_exact": True,
                        "preview_formula_cells": sum(
                            1
                            for row in [header, *preview_rows]
                            for value in row
                            if isinstance(value, dict) and value.get("formula") is True
                        ),
                    }
                )
    finally:
        workbook.close()
    return {"kind": "workbook", "sheets": summaries}


def _xlsx_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    try:
        root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    return ["".join(node.itertext())[:MAX_JSON_STRING] for node in root]


def _xlsx_sheet(
    archive: zipfile.ZipFile,
    name: str,
    shared: list[str],
    max_rows: int = MAX_SAMPLE_ROWS + 1,
) -> dict[str, Any]:
    root = ElementTree.fromstring(archive.read(name))
    rows: list[list[str]] = []
    for row in root.iter():
        if row.tag.rsplit("}", 1)[-1] != "row":
            continue
        values: list[str] = []
        for cell in list(row)[:MAX_XLSX_COLUMNS]:
            if cell.tag.rsplit("}", 1)[-1] != "c":
                continue
            value = next(
                (child.text or "" for child in cell if child.tag.rsplit("}", 1)[-1] == "v"), ""
            )
            if cell.attrib.get("t") == "s" and value.isdigit() and int(value) < len(shared):
                value = shared[int(value)]
            values.append(value[:MAX_JSON_STRING])
        if values:
            rows.append(values)
        if len(rows) >= max_rows:
            break
    return {"sheet": name.rsplit("/", 1)[-1], "columns": rows[0] if rows else [], "rows": rows[1:]}


def read_rows(root: Path, source_ref: str) -> list[dict[str, Any]]:
    """Read bounded records only after the caller has supplied authorized refs."""
    path = _resolve_ref(root, source_ref)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        sample = decode_text(_read_prefix(path, 256 * 1024)).splitlines()[:20]
        try:
            delimiter = csv.Sniffer().sniff("\n".join(sample[:10]), delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
        with path.open("rb") as binary:
            text_stream = io.TextIOWrapper(
                binary, encoding=detect_encoding(_read_prefix(path, 64 * 1024)), errors="strict"
            )
            rows = csv.DictReader(text_stream, delimiter=delimiter)
            return [
                {str(key).strip(): value for key, value in row.items() if key is not None}
                for row in itertools.islice(rows, MAX_NORMALIZE_ROWS)
            ]
    if suffix == ".json":
        records: list[dict[str, Any]] = []
        with path.open("rb") as stream:
            for prefix in ("item", "rows.item", "orders.item", "records.item", "data.item"):
                stream.seek(0)
                try:
                    for value in ijson.items(stream, prefix):
                        if isinstance(value, dict):
                            records.append(value)
                            if len(records) >= MAX_NORMALIZE_ROWS:
                                break
                    if records:
                        break
                except ijson.common.IncompleteJSONError as error:
                    raise ValueError("invalid_json") from error
        return records
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    records: list[dict[str, Any]] = []
    try:
        for worksheet in list(workbook.worksheets)[:MAX_XLSX_SHEETS]:
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            columns = [str(value or "").strip() for value in header[:MAX_XLSX_COLUMNS]]
            for values in rows:
                if not any(value not in (None, "") for value in values):
                    continue
                record = {
                    column: values[index] if index < len(values) else ""
                    for index, column in enumerate(columns)
                    if column
                }
                qualified = {
                    f"{worksheet.title}::{column}": value for column, value in record.items()
                }
                records.append({**qualified, **record, "__sheet_name": worksheet.title})
                if len(records) >= MAX_NORMALIZE_ROWS:
                    return records
    finally:
        workbook.close()
    return records


def read_json_document(root: Path, source_ref: str) -> dict[str, Any]:
    """Read one bounded JSON document from an authorized Workspace source."""
    path = _resolve_ref(root, source_ref)
    if path.suffix.lower() != ".json":
        raise ValueError("source_ref_must_be_json")
    if path.stat().st_size > 8 * 1024 * 1024:
        raise ValueError("json_source_exceeds_size_limit")
    with path.open("r", encoding=detect_encoding(_read_prefix(path, 64 * 1024))) as stream:
        payload = json.load(stream)
    if not isinstance(payload, dict):
        raise ValueError("json_source_must_be_object")
    return payload


def flatten_record(record: dict[str, Any]) -> dict[str, Any]:
    flattened: dict[str, Any] = {}

    def walk(value: Any, prefix: str) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                walk(child, f"{prefix}.{key}" if prefix else str(key))
        elif isinstance(value, list):
            return
        else:
            flattened[prefix] = value
            flattened.setdefault(prefix.rsplit(".", 1)[-1], value)

    walk(record, "")
    return flattened


def _read_prefix(path: Path, limit: int) -> bytes:
    with path.open("rb") as stream:
        raw = stream.read(limit + 1)
    if len(raw) > limit:
        return raw[:limit]
    return raw


def detect_encoding(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "utf-16-le", "utf-16-be"):
        try:
            raw.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    best = from_bytes(raw).best()
    if best is not None and best.encoding and best.percent_chaos < 10:
        return best.encoding
    raise ValueError("unsupported_text_encoding")


def decode_text(raw: bytes) -> str:
    return raw.decode(detect_encoding(raw))


def _iter_json_events(stream: BinaryIO):
    try:
        yield from ijson.parse(stream)
    except (ijson.common.IncompleteJSONError, ijson.common.JSONError) as error:
        raise ValueError("invalid_json") from error


def _json_depth(prefix: str) -> int:
    return sum(1 for token in prefix.split(".") if token and token != "item")


def _validate_json_stream(stream: BinaryIO) -> None:
    count = 0
    max_depth = 0
    try:
        for prefix, event, value in ijson.parse(stream):
            count += 1
            if count > MAX_JSON_NODES:
                raise ValueError("json_node_limit_exceeded")
            depth = prefix.count(".") + prefix.count("item")
            max_depth = max(max_depth, depth)
            if max_depth > MAX_JSON_DEPTH:
                raise ValueError("json_depth_limit_exceeded")
            if event == "string" and isinstance(value, str) and len(value) > MAX_JSON_STRING:
                raise ValueError("json_string_limit_exceeded")
    except (ijson.common.IncompleteJSONError, ijson.common.JSONError) as error:
        raise ValueError("invalid_json") from error


def propose_mapping(
    profiles: list[dict[str, Any]], requirement_profile: dict[str, Any]
) -> dict[str, Any]:
    fields: list[tuple[str, str, Any]] = []
    seen_fields: set[tuple[str, str]] = set()

    def add_field(entity_name: str, field_name: Any, unit: Any = None) -> None:
        entity = str(entity_name).strip()
        field = str(field_name).strip()
        identity = (entity, field)
        if entity and field and identity not in seen_fields:
            seen_fields.add(identity)
            fields.append((entity, field, unit))

    def add_entity(entity_name: str, entity: Any) -> None:
        if not isinstance(entity, dict):
            return
        declared = entity.get("fields")
        if isinstance(declared, dict):
            for field_name, spec in declared.items():
                add_field(
                    entity_name,
                    field_name,
                    spec.get("unit") if isinstance(spec, dict) else None,
                )
        elif isinstance(declared, list):
            for field in declared:
                if isinstance(field, dict):
                    add_field(
                        entity_name,
                        field.get("name", field.get("displayName", "")),
                        field.get("unit"),
                    )
                else:
                    add_field(entity_name, field)
        for key in ("required_fields", "requiredFields", "optional_fields", "optionalFields"):
            declared = entity.get(key, [])
            if isinstance(declared, list):
                for field in declared:
                    if isinstance(field, dict):
                        add_field(entity_name, field.get("name", field.get("displayName", "")))
                    else:
                        add_field(entity_name, field)

    entities = requirement_profile.get("entities")
    if entities is None:
        entities = requirement_profile.get("requiredEntities")
    if isinstance(entities, list):
        for entity in entities:
            if isinstance(entity, dict):
                add_entity(
                    entity.get("name")
                    or entity.get("displayName")
                    or entity.get("entity")
                    or entity.get("entity_type")
                    or "",
                    entity,
                )
    elif isinstance(entities, dict):
        for entity_name, entity in entities.items():
            add_entity(entity_name, entity)
    else:
        # The current Network Agent contract uses an entity-oriented map at the
        # top level. Only values that look like field declarations are accepted;
        # metadata such as country and analysis settings is ignored.
        for entity_name, entity in requirement_profile.items():
            if isinstance(entity, dict) and any(
                key in entity
                for key in ("fields", "required_fields", "requiredFields", "optional_fields", "optionalFields")
            ):
                add_entity(entity_name, entity)
    ranked: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for profile in profiles:
        structure = profile.get("structure", {})
        names = _field_names(structure)
        for entity, target, target_unit in fields:
            for source in names:
                score = similarity(source, target)
                if score < 0.80:
                    # Substring matches such as ``date`` -> ``candidate`` are
                    # useful search hints but are not safe mapping candidates.
                    continue
                display_name = str(profile.get("display_name", "")).lower()
                sheet_context = " ".join(
                    str(sheet.get("sheet", "")).lower() for sheet in structure.get("sheets", [])
                )
                context = f"{display_name} {sheet_context}"
                entity_key = "".join(char.lower() for char in entity if char.isalnum())
                entity_tokens = {
                    "city": ("city", "cities", "城市"),
                    "demandlocation": ("location", "locations", "demand_location", "需求点"),
                    "demandpoints": ("demand", "order", "orders", "需求", "location", "locations"),
                    "demand": ("demand", "order", "orders", "需求"),
                    "facility": ("facility", "facilities", "warehouse", "仓", "设施"),
                    "existingwarehouses": ("warehouse", "warehouses", "facility", "facilities", "仓", "设施"),
                    "warehouses": ("warehouse", "warehouses", "facility", "facilities", "仓", "设施"),
                    "assignment": ("assignment", "assignments", "allocation", "分配"),
                    "rate": ("rate", "rates", "cost", "费率", "成本"),
                    "route": ("route", "routes", "travel", "路线"),
                }.get(entity_key, ())
                known_entity_tokens = {
                    token
                    for values in {
                        "city": ("city", "cities", "城市"),
                        "demandlocation": ("location", "locations", "demand_location", "需求点"),
                        "demandpoints": ("demand", "order", "orders", "需求", "location", "locations"),
                        "demand": ("demand", "order", "orders", "需求"),
                        "facility": ("facility", "facilities", "warehouse", "仓", "设施"),
                        "existingwarehouses": ("warehouse", "warehouses", "facility", "facilities", "仓", "设施"),
                        "warehouses": ("warehouse", "warehouses", "facility", "facilities", "仓", "设施"),
                        "assignment": ("assignment", "assignments", "allocation", "分配"),
                        "rate": ("rate", "rates", "cost", "费率", "成本"),
                        "route": ("route", "routes", "travel", "路线"),
                    }.values()
                    for token in values
                }
                context_has_known_entity = any(token in context for token in known_entity_tokens)
                if context_has_known_entity and not any(
                    token in context for token in entity_tokens
                ):
                    continue
                if any(token in context for token in entity_tokens):
                    score = min(1.0, score + 0.08)
                ranked.setdefault((entity, target), []).append(
                    {
                        "source_ref": profile["source_ref"],
                        "source_asset_id": profile.get("source_asset_id"),
                        "source_field": source,
                        "target_entity": entity,
                        "target_field": target,
                        "source_unit": None,
                        "target_unit": target_unit,
                        "transformation": "identity",
                        "confidence": round(score, 3),
                        "requires_confirmation": True,
                        "reason": (
                            "deterministic alias/type candidate; confirm the complete "
                            "mapping revision"
                        ),
                    }
                )

    candidates: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    for target, options in ranked.items():
        options.sort(
            key=lambda item: (
                -float(item["confidence"]),
                str(item.get("source_ref", "")),
                str(item.get("source_field", "")),
            )
        )
        selected = options[0]
        candidates.append(selected)
        if (
            len(options) > 1
            and float(selected["confidence"]) - float(options[1]["confidence"]) < 0.15
        ):
            conflicts.append(
                {
                    "target_entity": target[0],
                    "target_field": target[1],
                    "candidates": options[:5],
                    "reason": (
                        "multiple source fields have indistinguishable confidence; "
                        "choose one before confirming"
                    ),
                }
            )
    return {
        "schema": "mapping_proposal.v1",
        "candidates": candidates,
        "conflicts": conflicts,
        "requires_confirmation": True,
    }


def _field_names(structure: dict[str, Any]) -> list[str]:
    if structure.get("kind") == "table":
        return [str(value) for value in structure.get("columns", []) if value]
    if structure.get("kind") == "workbook":
        sheets = structure.get("sheets", [])
        counts: dict[str, int] = {}
        for sheet in sheets:
            for value in sheet.get("columns", []):
                if value:
                    key = str(value).strip()
                    counts[key] = counts.get(key, 0) + 1
        fields: list[str] = []
        for sheet in sheets:
            sheet_name = str(sheet.get("sheet", "Sheet"))
            for value in sheet.get("columns", []):
                if not value:
                    continue
                key = str(value).strip()
                fields.append(f"{sheet_name}::{key}" if counts.get(key, 0) > 1 else key)
        return fields
    names: list[str] = []
    for array in structure.get("arrays", []):
        preview_rows = array.get("preview", {}).get("rows", [])
        for item in preview_rows:
            if item.get("kind") == "object":
                names.extend(item.get("fields", {}).keys())
    return names


def similarity(source: str, target: str) -> float:
    def normalize(value: str) -> str:
        return "".join(char.lower() for char in value if char.isalnum())

    left, right = normalize(source), normalize(target)
    if not left or not right:
        return 0.0
    if left == right:
        return 1.0
    if left in right or right in left:
        return 0.72
    aliases = {
        "qty": "quantity",
        "数量": "quantity",
        "件数": "quantity",
        "lat": "latitude",
        "纬度": "latitude",
        "lon": "longitude",
        "lng": "longitude",
        "经度": "longitude",
        "warehouse": "facility",
        "仓库": "facility",
        "仓": "facility",
        "dc": "facility",
        "仓库编号": "facility_id",
        "仓库id": "facility_id",
        "候选仓": "candidate",
        "区域": "region",
        "日期": "date",
        "时间": "date",
        "容量": "capacity",
        "成本": "cost",
    }
    if aliases.get(left) == right or aliases.get(right) == left:
        return 0.82
    semantic_aliases = {
        "需求点": {"demandlocationid", "demandlocation", "origin", "demandid"},
        "需求地点": {"demandlocationid", "demandlocation", "origin"},
        "需求编号": {"demandid", "demandlocationid"},
        "设施编号": {"facilityid"},
        "候选点": {"facilityid", "candidate"},
        "现有候选": {"existingorcandidate"},
        "固定成本": {"fixedcost"},
        "开启成本": {"openingcost"},
        "运输时间": {"traveltimehours", "transitdays"},
    }
    if right in semantic_aliases.get(left, set()) or left in semantic_aliases.get(right, set()):
        return 0.82
    return 0.0
