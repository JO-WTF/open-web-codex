from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from supply_chain_planner import workspace_intake
from supply_chain_planner.data_server import _build_planning_source
from supply_chain_planner.workspace_intake import discover, inspect, read_rows


def test_workspace_discovery_returns_bounded_relative_descriptors(tmp_path: Path) -> None:
    (tmp_path / "demand.csv").write_text(
        "origin,quantity,date\nJakarta,10,2026-01-01\n",
        encoding="utf-8",
    )
    (tmp_path / "nested").mkdir()
    (tmp_path / "nested" / "facilities.json").write_text(
        json.dumps({"facilities": [{"facility_id": "JKT", "latitude": -6.2}]}),
        encoding="utf-8",
    )
    (tmp_path / "notes.txt").write_text("ignored", encoding="utf-8")

    sources = discover(tmp_path)

    assert [source["relative_path"] for source in sources] == [
        "demand.csv",
        "nested/facilities.json",
    ]
    assert sources == [
        {
            "relative_path": "demand.csv",
            "format": "csv",
            "size": (tmp_path / "demand.csv").stat().st_size,
        },
        {
            "relative_path": "nested/facilities.json",
            "format": "json",
            "size": (tmp_path / "nested" / "facilities.json").stat().st_size,
        },
    ]


def test_workspace_discovery_rejects_legacy_excel_formats(tmp_path: Path) -> None:
    (tmp_path / "legacy.xlsm").write_bytes(b"not a supported workbook")

    with pytest.raises(ValueError, match="unsupported_source_format"):
        discover(tmp_path)


def test_exact_workspace_source_validation_rejects_escape_symlink_and_size(
    tmp_path: Path, monkeypatch
) -> None:
    (tmp_path / "valid.csv").write_text("id\n1\n", encoding="utf-8")
    (tmp_path / "empty.csv").touch()
    (tmp_path / "unsupported.txt").write_text("id\n1\n", encoding="utf-8")
    outside = tmp_path.parent / f"{tmp_path.name}-outside.csv"
    outside.write_text("id\n2\n", encoding="utf-8")
    (tmp_path / "linked.csv").symlink_to(outside)
    (tmp_path / "linked-dir").symlink_to(tmp_path.parent, target_is_directory=True)

    for relative_path, expected in [
        ("", "relative_path_required"),
        (".", "invalid_workspace_relative_path"),
        ("../outside.csv", "invalid_workspace_relative_path"),
        (str((tmp_path / "valid.csv").resolve()), "invalid_workspace_relative_path"),
        ("linked.csv", "symlink_rejected"),
        (f"linked-dir/{outside.name}", "symlink_rejected"),
        ("unsupported.txt", "unsupported_source_format"),
        ("empty.csv", "source_size_limit"),
    ]:
        with pytest.raises(ValueError, match=expected):
            inspect(tmp_path, relative_path)

    monkeypatch.setattr(workspace_intake, "MAX_BYTES", 4)
    with pytest.raises(ValueError, match="source_size_limit"):
        inspect(tmp_path, "valid.csv")


def test_source_structure_limits_reject_instead_of_truncating(tmp_path: Path, monkeypatch) -> None:
    (tmp_path / "rows.csv").write_text("id\n1\n2\n", encoding="utf-8")
    monkeypatch.setattr(workspace_intake, "MAX_NORMALIZE_ROWS", 1)
    with pytest.raises(ValueError, match="row_limit"):
        inspect(tmp_path, "rows.csv")

    workbook = Workbook()
    workbook.active.append(["first", "second"])
    workbook.save(tmp_path / "columns.xlsx")
    monkeypatch.setattr(workspace_intake, "MAX_XLSX_COLUMNS", 1)
    with pytest.raises(ValueError, match="column_limit"):
        inspect(tmp_path, "columns.xlsx")


def test_csv_and_json_profiles_are_structural_samples_with_exact_counts(tmp_path: Path) -> None:
    csv_path = tmp_path / "demand.csv"
    csv_path.write_text(
        "origin;qty\n" + "\n".join(f"city-{index};{index}" for index in range(25)) + "\n",
        encoding="utf-8",
    )
    json_path = tmp_path / "nested.json"
    json_path.write_text(
        json.dumps(
            {"orders": [{"origin": f"city-{index}", "quantity": index} for index in range(25)]}
        ),
        encoding="utf-8",
    )

    sources = discover(tmp_path)
    profiles = {
        source["relative_path"]: inspect(tmp_path, source["relative_path"]) for source in sources
    }

    assert profiles["demand.csv"]["structure"]["delimiter"] == ";"
    assert profiles["demand.csv"]["structure"]["columns"] == ["origin", "qty"]
    csv_structure = profiles["demand.csv"]["structure"]
    assert csv_structure["record_count"] == 25
    assert csv_structure["record_count_exact"] is True
    assert csv_structure["preview"]["strategy"] == "head"
    assert csv_structure["preview"]["returned_count"] == 20
    assert csv_structure["preview"]["complete"] is False
    assert profiles["nested.json"]["structure"]["arrays"][0]["path"] == "$.orders"
    json_array = profiles["nested.json"]["structure"]["arrays"][0]
    assert json_array["length"] == 25
    assert json_array["length_exact"] is True
    assert json_array["preview"]["returned_count"] == 3
    assert json_array["preview"]["complete"] is False


def test_read_rows_supports_explicit_rows_array(tmp_path: Path) -> None:
    (tmp_path / "administrative-areas.json").write_text(
        json.dumps({"rows": [{"city_id": "IDN-CITY-001", "city_name": "Jakarta"}]}),
        encoding="utf-8",
    )
    source = discover(tmp_path)[0]

    assert read_rows(tmp_path, source["relative_path"]) == [
        {"city_id": "IDN-CITY-001", "city_name": "Jakarta"}
    ]


def test_xlsx_profile_preserves_multiple_sheets_and_bounded_rows(tmp_path: Path) -> None:
    workbook = Workbook()
    locations = workbook.active
    locations.title = "需求点"
    locations.append(["Demand Location ID", "Name", "Region"])
    locations.append(["D1", "Jakarta", "ID-JK"])
    facilities = workbook.create_sheet("仓库")
    facilities.append(["Facility ID", "Name", "Existing or Candidate"])
    facilities.append(["F1", "Existing", "existing"])
    workbook.save(tmp_path / "network.xlsx")

    source = discover(tmp_path)[0]
    profile = inspect(tmp_path, source["relative_path"])

    assert source["format"] == "xlsx"
    assert [sheet["sheet"] for sheet in profile["structure"]["sheets"]] == ["需求点", "仓库"]
    assert profile["structure"]["sheets"][0]["columns"] == [
        "Demand Location ID",
        "Name",
        "Region",
    ]
    warehouse_sheet = profile["structure"]["sheets"][1]
    assert warehouse_sheet["preview"]["rows"] == [["F1", "Existing", "existing"]]
    assert warehouse_sheet["record_count"] == 1
    assert warehouse_sheet["record_count_exact"] is True


def test_confirmed_mapping_builds_strict_planning_source(tmp_path: Path) -> None:
    files = {
        "cities.csv": (
            "city_id,name,region,latitude,longitude\n"
            "CITY-JKT,Jakarta,ID-JK,-6.2,106.8\n"
            "CITY-BDG,Bandung,ID-JB,-6.9,107.6\n"
        ),
        "city-demand.csv": "city_id,date,quantity\nCITY-JKT,2026-01-01,10\n",
        "facilities.csv": (
            "facility_id,name,city_id,existing_or_candidate,capacity,"
            "fixed_cost,opening_cost,handling_cost\n"
            "F1,Existing,CITY-JKT,existing,100,100,0,1\n"
            "F2,Candidate,CITY-BDG,candidate,100,100,20,1\n"
        ),
        "coverage.csv": "facility_id,city_id,is_current\nF1,CITY-JKT,true\n",
        "lanes.csv": (
            "origin_city_id,destination_city_id,distance_km,travel_time_hours,"
            "base_cost_per_unit,distance_cost_per_km_per_unit,currency\n"
            "CITY-JKT,CITY-JKT,0,0,1,0.1,IDR\n"
            "CITY-BDG,CITY-JKT,150,4,1,0.1,IDR\n"
        ),
    }
    for name, content in files.items():
        (tmp_path / name).write_text(content, encoding="utf-8")

    sources = discover(tmp_path)
    refs = {source["relative_path"]: source["relative_path"] for source in sources}
    mapping = [
        {
            "relative_path": refs["cities.csv"],
            "source_field": field,
            "target_entity": "City",
            "target_field": field,
        }
        for field in ("city_id", "name", "region", "latitude", "longitude")
    ]
    mapping += [
        {
            "relative_path": refs["city-demand.csv"],
            "source_field": field,
            "target_entity": "CityDemand",
            "target_field": field,
        }
        for field in ("city_id", "date", "quantity")
    ]
    mapping += [
        {
            "relative_path": refs["facilities.csv"],
            "source_field": field,
            "target_entity": "Facility",
            "target_field": field,
        }
        for field in (
            "facility_id",
            "name",
            "city_id",
            "existing_or_candidate",
            "capacity",
            "fixed_cost",
            "opening_cost",
            "handling_cost",
        )
    ]
    mapping += [
        {
            "relative_path": refs["coverage.csv"],
            "source_field": field,
            "target_entity": "Coverage",
            "target_field": field,
        }
        for field in ("facility_id", "city_id", "is_current")
    ]
    mapping += [
        {
            "relative_path": refs["lanes.csv"],
            "source_field": field,
            "target_entity": "Lane",
            "target_field": field,
        }
        for field in (
            "origin_city_id",
            "destination_city_id",
            "distance_km",
            "travel_time_hours",
            "base_cost_per_unit",
            "distance_cost_per_km_per_unit",
            "currency",
        )
    ]
    source = _build_planning_source(
        tmp_path,
        list(refs.values()),
        mapping,
        [
            {"name": "market", "value": "ID"},
            {"name": "planning_mode", "value": "candidate_warehouse_optimization"},
            {"name": "cost_scope", "value": "fixed, opening, handling, transport"},
            {"name": "target_sla_hours", "value": 48},
            {"name": "coverage_target", "value": 95},
        ],
    )

    assert source.currency == "IDR"
    assert source.planning_period == "2026-01"
    assert source.route_method == "quoted"
    assert source.route_provider == "workspace-quoted-lanes"
    assert len(source.city_demands) == 1
    assert len(source.warehouse_city_coverage) == 1
    assert len(source.lanes) == 2
